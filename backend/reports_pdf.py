import base64, io
from datetime import datetime, timedelta, timezone
from pathlib import Path
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ASSETS = Path(__file__).parent / "assets"
WIB = timezone(timedelta(hours=7))
NAVY = colors.HexColor("#0f2745")
BLUE = colors.HexColor("#0284c7")
LIGHT = colors.HexColor("#f1f5f9")

TITLES = {"transaksi": "Laporan Transaksi Peminjaman", "masuk": "Laporan Buku Masuk", "hilang": "Laporan Buku Hilang", "kunjungan": "Laporan Pengunjung Perpustakaan", "statistik": "Laporan Statistik Pengunjung", "peringkat": "Peringkat Rajin Berkunjung"}
PERIODS = {"semua": "Semua waktu", "harian": "Harian", "bulanan": "Bulanan", "7hari": "7 hari terakhir", "30hari": "30 hari terakhir"}
CHECKOUT_LABEL = {"scan": "Scan", "manual": "Petugas", "otomatis": "Otomatis"}


BULAN = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]
BULAN_PANJANG = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]


def fmt_long(dt, with_time=False):
    base = f"{dt.day} {BULAN_PANJANG[dt.month - 1]} {dt.year}"
    return f"{base}, {dt.strftime('%H:%M')}" if with_time else base


def fmt_date(value, with_time=False):
    if not value: return "-"
    try:
        dt = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return value
    if dt.tzinfo: dt = dt.astimezone(WIB)
    base = f"{dt.day:02d} {BULAN[dt.month - 1]} {dt.year}"
    return f"{base} {dt.strftime('%H:%M')}" if with_time else base


def in_period(value, period, now):
    if period == "semua" or not value: return True
    try:
        dt = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return True
    dt = dt.astimezone(WIB) if dt.tzinfo else dt
    if period == "harian": return dt.date() == now.date()
    if period == "bulanan": return (dt.year, dt.month) == (now.year, now.month)
    return True


def build_loans(transactions):
    loans, open_by_book = [], {}
    for t in sorted(transactions, key=lambda x: x["date"]):
        if t["action"] == "pinjam":
            loan = {"id": t["id"], "member_id": t["member_id"], "member_name": t["member_name"], "class_name": t.get("class_name", ""), "book_id": t["book_id"], "book_title": t["book_title"], "borrowed_at": t["date"], "returned_at": None, "status": "Dipinjam", "date": t["date"]}
            loans.append(loan); open_by_book.setdefault(t["book_id"], []).append(loan)
        else:
            stack = open_by_book.get(t["book_id"]) or []
            if stack:
                idx = next((i for i, l in enumerate(stack) if l["member_id"] == t["member_id"]), 0)
                loan = stack.pop(idx); loan.update(returned_at=t["date"], status="Dikembalikan")
            else:
                loans.append({"id": t["id"], "member_id": t["member_id"], "member_name": t["member_name"], "class_name": t.get("class_name", ""), "book_id": t["book_id"], "book_title": t["book_title"], "borrowed_at": None, "returned_at": t["date"], "status": "Dikembalikan", "date": t["date"]})
    loans.sort(key=lambda l: l["returned_at"] or l["borrowed_at"] or "", reverse=True)
    return loans


def table_rows(kind, rows):
    if kind == "transaksi":
        header = ["No", "Anggota", "Kelas", "Buku", "Kode", "Tgl Pinjam", "Tgl Kembali", "Status"]
        widths = [10 * mm, 44 * mm, 24 * mm, 62 * mm, 24 * mm, 34 * mm, 34 * mm, 26 * mm]
        plain = [[r["member_name"], r.get("class_name") or "-", r["book_title"], r["book_id"], fmt_date(r["borrowed_at"], True), fmt_date(r["returned_at"], True), r["status"]] for r in rows]
        borrowed = sum(1 for r in rows if r["status"] == "Dipinjam"); summary = f"Total peminjaman: {len(rows)} · Masih dipinjam: {borrowed} · Dikembalikan: {len(rows) - borrowed}"
    elif kind == "masuk":
        header = ["No", "Judul Buku", "Penulis", "Jumlah", "Sumber", "Rak", "Tanggal", "Dicatat oleh"]
        widths = [10 * mm, 70 * mm, 45 * mm, 18 * mm, 32 * mm, 16 * mm, 30 * mm, 36 * mm]
        plain = [[r["title"], r.get("author", "-"), f"{r['quantity']} eks", r.get("source") or "-", r.get("shelf", "-"), fmt_date(r["date"]), r.get("recorded_by", "-")] for r in rows]
        summary = f"Total penerimaan: {len(rows)} · Total eksemplar: {sum(r.get('quantity', 0) for r in rows)}"
    elif kind == "kunjungan":
        header = ["No", "Nama", "Jenis", "Kelas / Asal", "Keperluan", "Masuk", "Keluar", "Cara keluar"]
        widths = [10 * mm, 52 * mm, 20 * mm, 36 * mm, 46 * mm, 34 * mm, 34 * mm, 28 * mm]
        plain = [[r["name"], "Anggota" if r.get("kind") == "anggota" else "Tamu", (r.get("class_name") if r.get("kind") == "anggota" else (r.get("origin") or "Tamu")) or "-", r.get("purpose") or "-", fmt_date(r.get("check_in"), True), fmt_date(r.get("check_out"), True) if r.get("check_out") else "Masih di dalam", CHECKOUT_LABEL.get(r.get("checkout_type"), "-")] for r in rows]
        mem = sum(1 for r in rows if r.get("kind") == "anggota"); n = len(rows) or 1
        summary = f"Total pengunjung: {len(rows)} · Anggota: {mem} ({mem * 100 / n:.1f}%) · Tamu: {len(rows) - mem} ({(len(rows) - mem) * 100 / n:.1f}%)" if rows else "Total pengunjung: 0"
    elif kind == "statistik":
        st = rows; header = ["No", "Kelas", "Jumlah pengunjung", "Persentase"]
        widths = [12 * mm, 120 * mm, 60 * mm, 60 * mm]
        plain = [[c["class_name"], str(c["count"]), f"{c['pct']:.1f}%"] for c in st["per_class"]]
        kinds = " · ".join(f"{k['label']}: {k['count']} ({k['pct']:.1f}%)" for k in st["kinds"])
        summary = f"Total pengunjung: {st['total']} · Rata-rata {st['avg_per_day']}/hari · {kinds} · Durasi rata-rata {st['avg_duration_min']} menit"
    elif kind == "peringkat":
        lb = rows; header = ["Peringkat", "Nama Siswa", "Kelas", "Pengunjung", "Hari datang", "Total waktu", "Persentase"]
        widths = [22 * mm, 70 * mm, 40 * mm, 30 * mm, 30 * mm, 34 * mm, 30 * mm]
        plain = [[s_["name"], s_["class_name"], str(s_["visits"]), str(s_["days"]), f"{s_.get('minutes', 0) // 60} j {s_.get('minutes', 0) % 60} m", f"{s_['pct']:.1f}%"] for s_ in lb["students"]]
        top_cls = lb["classes"][0]["class_name"] if lb["classes"] else "-"
        summary = f"{lb['label']} · Total pengunjung anggota: {lb['total_visits']} · Siswa aktif: {lb['active_students']} · Kelas teraktif: {top_cls}"
    else:
        header = ["No", "Judul Buku", "Kode", "Catatan", "Tgl Lapor", "Status", "Ditemukan", "Pelapor"]
        widths = [10 * mm, 62 * mm, 24 * mm, 56 * mm, 30 * mm, 22 * mm, 30 * mm, 30 * mm]
        plain = [[r["title"], r["code"], r.get("detail") if r.get("detail") != r.get("code") else "-", fmt_date(r["date"]), r["status"], fmt_date(r.get("found_at")), r.get("reported_by", "-")] for r in rows]
        lost = sum(1 for r in rows if r["status"] == "Hilang"); summary = f"Total laporan: {len(rows)} · Masih hilang: {lost} · Ditemukan: {len(rows) - lost}"
    return header, widths, plain, summary


def render_report_xlsx(kind, period, rows, settings, officer, orientation="landscape"):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    now = datetime.now(WIB)
    header, widths, plain, summary = table_rows(kind, rows)
    wb = Workbook(); ws = wb.active; ws.title = TITLES[kind][:30]
    ncol = len(header)
    for logo_val, fallback, col in ((settings.get("school_logo"), "logo-sekolah.jpg", "A"), (settings.get("library_logo"), "logo-perpustakaan.jpg", get_column_letter(ncol))):
        try:
            img = _logo_bytes(logo_val, fallback); xl = XLImage(img); xl.width = xl.height = 60; ws.add_image(xl, f"{col}1")
        except Exception:
            pass
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=ncol - 1); ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=ncol - 1); ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=ncol - 1)
    ws["B1"] = "PERPUSTAKAAN SMA NEGERI 1 ABUNG TENGAH"; ws["B1"].font = Font(bold=True, size=14, color="0F2745"); ws["B1"].alignment = Alignment(horizontal="center")
    ws["B2"] = "Jln Raya Subik, Desa Pekurun Selatan, Lampung Utara"; ws["B2"].alignment = Alignment(horizontal="center"); ws["B2"].font = Font(size=9, color="475569")
    ws["B3"] = f"{TITLES[kind]} — {PERIODS.get(period, period or 'Semua waktu')} · dicetak {fmt_long(now, True)} WIB oleh {officer}"; ws["B3"].alignment = Alignment(horizontal="center"); ws["B3"].font = Font(size=9, italic=True, color="475569")
    ws.row_dimensions[1].height = 24; ws.row_dimensions[2].height = 14; ws.row_dimensions[4].height = 6
    ws.cell(row=5, column=1, value=summary).font = Font(bold=True, color="0284C7")
    thin = Side(style="thin", color="CBD5E1"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for j, h in enumerate(header, 1):
        c = ws.cell(row=7, column=j, value=h); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F2745"); c.alignment = Alignment(horizontal="center", vertical="center"); c.border = border
    for i, r in enumerate(plain, 1):
        for j, v in enumerate([i] + r, 1):
            c = ws.cell(row=7 + i, column=j, value=v); c.border = border; c.alignment = Alignment(vertical="center", wrap_text=True)
            if i % 2 == 0: c.fill = PatternFill("solid", fgColor="F1F5F9")
    if not plain: ws.cell(row=8, column=2, value="Belum ada data pada periode ini")
    if kind == "peringkat" and rows["classes"]:
        start = 7 + max(len(plain), 1) + 3
        ws.cell(row=start - 1, column=1, value="Peringkat kelas").font = Font(bold=True, color="0F2745")
        for j, h in enumerate(["Peringkat", "Kelas", "Pengunjung", "Siswa berkunjung", "Persentase"], 1):
            c = ws.cell(row=start, column=j, value=h); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F2745"); c.border = border
        for i, cl in enumerate(rows["classes"], 1):
            for j, v in enumerate([cl["rank"], cl["class_name"], cl["visits"], cl["members"], f"{cl['pct']:.1f}%"], 1):
                c = ws.cell(row=start + i, column=j, value=v); c.border = border
    if kind == "statistik":
        start = 7 + max(len(plain), 1) + 3
        ws.cell(row=start - 1, column=1, value="Pengunjung per hari").font = Font(bold=True, color="0F2745")
        tot = rows["total"] or 1
        for j, h in enumerate(["Tanggal", "Total", "Anggota", "Tamu", "Persentase"], 1):
            c = ws.cell(row=start, column=j, value=h); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F2745"); c.border = border
        for i, d in enumerate(rows["per_day"], 1):
            for j, v in enumerate([fmt_date(d["date"]), d["total"], d["anggota"], d["tamu"], f"{d['total'] * 100 / tot:.1f}%"], 1):
                c = ws.cell(row=start + i, column=j, value=v); c.border = border
    for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = max(6, w / mm * 0.55)
    ws.freeze_panes = "A8"
    # Pengaturan cetak Excel mengikuti orientasi yang dipilih: A4, lebar pas 1 halaman, judul tabel diulang.
    ws.page_setup.orientation = "portrait" if orientation == "portrait" else "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "7:7"
    ws.page_margins.left = ws.page_margins.right = 0.4; ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _logo_bytes(value, fallback):
    from PIL import Image as PILImage
    if value and value.startswith("data:"):
        src = PILImage.open(io.BytesIO(base64.b64decode(value.split(",", 1)[1]))).convert("RGBA")
        bg = PILImage.new("RGBA", src.size, (255, 255, 255, 255)); bg.alpha_composite(src); src = bg.convert("RGB").resize((180, 180))
    else:
        src = PILImage.open(str(ASSETS / fallback)).convert("RGB")
    out = io.BytesIO(); src.save(out, format="PNG"); out.seek(0); return out


def _logo(value, fallback, size=16 * mm):
    try:
        if value and value.startswith("data:"):
            from PIL import Image as PILImage
            src = PILImage.open(io.BytesIO(base64.b64decode(value.split(",", 1)[1]))).convert("RGBA")
            bg = PILImage.new("RGBA", src.size, (255, 255, 255, 255)); bg.alpha_composite(src)
            out = io.BytesIO(); bg.convert("RGB").resize((180, 180)).save(out, format="JPEG", quality=88); out.seek(0)
            img = Image(out)
        else:
            img = Image(str(ASSETS / fallback))
        img.drawWidth = img.drawHeight = size
        return img
    except Exception:
        return Spacer(size, size)


def _fit(widths, avail):
    """Skalakan lebar kolom agar tabel pas selebar halaman (potret/lanskap)."""
    total = sum(widths) or 1
    return [w * avail / total for w in widths]


def _table_style():
    return TableStyle([("BACKGROUND", (0, 0), (-1, 0), NAVY), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)])


def render_report_pdf(kind, period, rows, settings, officer, orientation="landscape"):
    now = datetime.now(WIB)
    buf = io.BytesIO()
    portrait = orientation == "portrait"
    pagesize = A4 if portrait else landscape(A4)
    margin_x = 12 * mm if portrait else 14 * mm
    avail = pagesize[0] - 2 * margin_x
    doc = SimpleDocTemplate(buf, pagesize=pagesize, leftMargin=margin_x, rightMargin=margin_x, topMargin=12 * mm, bottomMargin=14 * mm, title=TITLES[kind])
    # Ukuran huruf menyesuaikan lebar kertas: potret lebih sempit -> huruf tabel sedikit lebih kecil & teks dibungkus.
    fs = 7.6 if portrait else 8.5
    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=13.5 if portrait else 15, textColor=NAVY, leading=17 if portrait else 18)
    sub = ParagraphStyle("sub", fontName="Helvetica", fontSize=8.5 if portrait else 9, textColor=colors.HexColor("#475569"), leading=11.5 if portrait else 12)
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=fs, leading=fs + 2.5, wordWrap="LTR", splitLongWords=True)
    cellb = ParagraphStyle("cellb", parent=cell, fontName="Helvetica-Bold")
    head = ParagraphStyle("head", fontName="Helvetica-Bold", fontSize=fs - 0.6, leading=fs + 1.5, textColor=colors.white)

    logo = 15 * mm if portrait else 16 * mm
    kop = Table([[_logo(settings.get("school_logo"), "logo-sekolah.jpg", logo), [Paragraph("PERPUSTAKAAN SMA NEGERI 1 ABUNG TENGAH", h1), Paragraph("Jln Raya Subik, Desa Pekurun Selatan, Lampung Utara · Tanpa denda, budaya membaca", sub)], _logo(settings.get("library_logo"), "logo-perpustakaan.jpg", logo)]], colWidths=[logo + 4 * mm, avail - 2 * (logo + 4 * mm), logo + 4 * mm])
    kop.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 0), (-1, 0), 1.5, NAVY), ("BOTTOMPADDING", (0, 0), (-1, 0), 6)]))

    title = Paragraph(f"{TITLES[kind]} — {PERIODS.get(period, period or 'Semua waktu')}", ParagraphStyle("t", parent=h1, fontSize=12 if portrait else 13, spaceBefore=8))
    meta = Paragraph(f"Dicetak {fmt_long(now, True)} WIB · {len(rows) if isinstance(rows, list) else rows.get('total', rows.get('total_visits', 0))} data · oleh {officer} · A4 {'potret' if portrait else 'lanskap'}", sub)

    header, widths, plain, summary = table_rows(kind, rows)
    body = [[Paragraph(str(i + 1), cell)] + [Paragraph(str(v), cellb if j == 0 else cell) for j, v in enumerate(r)] for i, r in enumerate(plain)]

    if not body: body = [["", Paragraph("Belum ada data pada periode ini", cell)] + [""] * (len(header) - 2)]
    table = Table([[Paragraph(h, head) for h in header]] + body, colWidths=_fit(widths, avail), repeatRows=1)
    table.setStyle(_table_style())

    def side_table(headers, data, ws):
        t = Table([[Paragraph(h, head) for h in headers]] + [[Paragraph(str(v), cell) for v in r] for r in data], colWidths=_fit(ws, avail), repeatRows=1)
        t.setStyle(_table_style()); return t

    extra = []
    if kind == "statistik":
        days = [d for d in rows["per_day"]][-31:]
        tot = rows["total"] or 1
        dt = side_table(["Tanggal", "Total", "Anggota", "Tamu", "Persentase"], [[fmt_date(d["date"]), d["total"], d["anggota"], d["tamu"], f"{d['total'] * 100 / tot:.1f}%"] for d in days], [60, 40, 40, 40, 40])
        extra = [Spacer(1, 12), Paragraph("Pengunjung per hari", ParagraphStyle("t2", parent=h1, fontSize=11)), Spacer(1, 4), dt]
    if kind == "peringkat" and rows["classes"]:
        ct = side_table(["Peringkat", "Kelas", "Pengunjung", "Siswa berkunjung", "Persentase"], [[c["rank"], c["class_name"], c["visits"], c["members"], f"{c['pct']:.1f}%"] for c in rows["classes"]], [30, 70, 40, 50, 40])
        extra = [Spacer(1, 12), Paragraph("Peringkat kelas", ParagraphStyle("t3", parent=h1, fontSize=11)), Spacer(1, 4), ct]
    sign = Table([["", Paragraph(f"Pekurun Selatan, {fmt_long(now)}<br/>Petugas Perpustakaan,<br/><br/><br/><br/><b>{officer}</b>", ParagraphStyle("s", parent=cell, fontSize=9, leading=12, alignment=1))]], colWidths=[avail - 70 * mm, 70 * mm])
    doc.build([kop, title, meta, Spacer(1, 4), Paragraph(summary, ParagraphStyle("sum", parent=sub, textColor=BLUE, fontName="Helvetica-Bold")), Spacer(1, 6), table, *extra, Spacer(1, 14), sign])
    return buf.getvalue()
