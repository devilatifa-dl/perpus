"""Iteration 9: xlsx-only member import + card back design settings."""
import os, io, pytest, requests
from openpyxl import Workbook, load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = BASE_URL + "/api"

ADMIN = {"email": "admin@smansa.sch.id", "password": "admin123", "portal": "admin"}
SISWA = {"email": "siswa@smansa.sch.id", "password": "siswa123", "portal": "siswa"}


def _login(creds):
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json=creds)
    assert r.status_code == 200, r.text
    return s


@pytest.fixture(scope="module")
def admin():
    return _login(ADMIN)


@pytest.fixture(scope="module")
def siswa():
    return _login(SISWA)


def _xlsx(rows, headers=("nama", "nis", "kelas")):
    wb = Workbook(); ws = wb.active
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.getvalue()


# --- Settings: card_back_* fields ---
class TestCardBackSettings:
    def test_defaults_present(self, admin):
        r = admin.get(f"{API}/settings"); assert r.status_code == 200
        d = r.json()
        assert "card_back_title" in d and d["card_back_title"]
        assert "card_back_rules" in d and "\n" in d["card_back_rules"]
        assert "card_back_web" in d

    def test_persist_update(self, admin):
        orig = admin.get(f"{API}/settings").json()
        payload = {**orig,
                   "card_back_title": "TEST BACK TITLE",
                   "card_back_rules": "Aturan 1\nAturan 2\nAturan 3",
                   "card_back_web": "test.example.sch.id"}
        r = admin.put(f"{API}/settings", json=payload); assert r.status_code == 200, r.text
        d = admin.get(f"{API}/settings").json()
        assert d["card_back_title"] == "TEST BACK TITLE"
        assert d["card_back_rules"] == "Aturan 1\nAturan 2\nAturan 3"
        assert d["card_back_web"] == "test.example.sch.id"
        # restore
        admin.put(f"{API}/settings", json=orig)

    def test_empty_web_persist(self, admin):
        orig = admin.get(f"{API}/settings").json()
        admin.put(f"{API}/settings", json={**orig, "card_back_web": ""})
        d = admin.get(f"{API}/settings").json()
        assert d["card_back_web"] == ""
        admin.put(f"{API}/settings", json=orig)


# --- Import template ---
class TestImportTemplate:
    def test_staff_download(self, admin):
        r = admin.get(f"{API}/members/import-template.xlsx")
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")
        wb = load_workbook(io.BytesIO(r.content), read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        assert [c.lower() if c else "" for c in rows[0][:3]] == ["nama", "nis", "kelas"]

    def test_siswa_forbidden(self, siswa):
        r = siswa.get(f"{API}/members/import-template.xlsx")
        assert r.status_code == 403


# --- Import members ---
class TestImport:
    def test_csv_rejected(self, admin):
        csv_bytes = b"nama,nis,kelas\nTEST_CsvAnggota,99991,XII-TEST\n"
        r = admin.post(f"{API}/members/import",
                       files={"file": ("members.csv", csv_bytes, "text/csv")})
        assert r.status_code == 400
        assert "Excel" in r.json().get("detail", "") and ".xlsx" in r.json().get("detail", "")

    def test_xlsx_import_numeric_nis(self, admin):
        # numeric cells shouldn't produce "99992.0"
        content = _xlsx([("TEST_ImpA", 99992, "XII-TEST"),
                         ("TEST_ImpB", 99993, "XII-TEST")])
        try:
            r = admin.post(f"{API}/members/import",
                           files={"file": ("m.xlsx", content,
                                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
            assert r.status_code == 200, r.text
            body = r.json()
            assert body["imported"] == 2
            assert body["total"] == 2
            # verify persisted with correct student_id string
            ms = admin.get(f"{API}/members").json()
            found = [m for m in ms if m["name"] in ("TEST_ImpA", "TEST_ImpB")]
            assert len(found) == 2
            for m in found:
                assert m["student_id"] in ("99992", "99993"), f"unexpected NIS: {m['student_id']}"
                assert "." not in m["student_id"]
        finally:
            ms = admin.get(f"{API}/members").json()
            for m in ms:
                if m["name"].startswith("TEST_Imp"):
                    admin.delete(f"{API}/members/{m['id']}")

    def test_siswa_import_forbidden(self, siswa):
        content = _xlsx([("TEST_X", 1, "X")])
        r = siswa.post(f"{API}/members/import",
                       files={"file": ("m.xlsx", content,
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert r.status_code == 403
