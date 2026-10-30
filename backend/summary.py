import io
from datetime import date, datetime, timedelta
from fastapi import APIRouter, Depends, HTTPException, Response
from reports_pdf import WIB, build_loans, fmt_long, _logo_bytes

MONTHS = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]


def wib_day(value):
    if not value:
        return None
    try:
        dt = datetime.fromisoformat(str(value))
    except ValueError:
        return None
    return dt.astimezone(WIB).date() if dt.tzinfo else dt.date()


def bucket(d, group):
    return d.strftime("%Y-%m") if group == "month" else d.isoformat()


def bucket_label(key, group):
    if group == "month":
        y, m = key.split("-")
        return f"{MONTHS[int(m) - 1]} {y}"
    d = date.fromisoformat(key)
    return f"{d.day} {MONTHS[d.month - 1]}"


def series_keys(d1, d2, group):
    keys, cur = [], d1
    while cur <= d2:
        k = bucket(cur, group)
        if not keys or keys[-1] != k:
            keys.append(k)
        cur += timedelta(days=1)
    return keys


def build_router(db, current_user, staff_only, overdue_loans, visits_auto_close):
    r = APIRouter()

    @r.get("/reports/summary")
    async def report_summary(start: str = "", end: str = "", group: str = "day", user=Depends(current_user)):
        staff_only(user)
        return await compute_summary(start, end, group)

    @r.get("/reports/summary.xlsx")
    async def report_summary_xlsx(start: str = "", end: str = "", group: str = "day", orientation: str = "landscape", user=Depends(current_user)):
        staff_only(user)
        d = await compute_summary(start, end, group)
        settings = await db.settings.find_one({"id": "main"}, {"_id": 0}) or {}
        data = render_summary_xlsx(d, settings, user["name"], "portrait" if orientation == "portrait" else "landscape")
        fname = f"laporan-ringkasan-{d['start']}-{d['end']}.xlsx"
        return Response(content=data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{fname}"'})

    @r.get("/reports/summary.pdf")
    async def report_summary_pdf(start: str = "", end: str = "", group: str = "day", orientation: str = "portrait", user=Depends(current_user)):
        staff_only(user)
        d = await compute_summary(start, end, group)
        settings = await db.settings.find_one({"id": "main"}, {"_id": 0}) or {}
        o = "landscape" if orientation == "landscape" else "portrait"
        data = render_summary_pdf(d, settings, user["name"], o)
        fname = f"laporan-ringkasan-{d['start']}-{d['end']}-{'lanskap' if o == 'landscape' else 'potret'}.pdf"
        return Response(content=data, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{fname}"'})

    async def compute_summary(start, end, group):
        if group not in ("day", "month"):
            raise HTTPException(400, "Pengelompokan tidak valid")
        today = datetime.now(WIB).date()
        try:
            d1 = date.fromisoformat(start) if start else today - timedelta(days=29)
            d2 = date.fromisoformat(end) if end else today
        except ValueError:
            raise HTTPException(400, "Format tanggal tidak valid")
        if d1 > d2:
            d1, d2 = d2, d1
        if (d2 - d1).days > 3 * 366:
            raise HTTPException(400, "Rentang tanggal maksimal 3 tahun")
        await visits_auto_close()

        loans = build_loans(await db.transactions.find({}, {"_id": 0}).to_list(100000))
        visits = await db.visits.find({"date": {"$gte": d1.isoformat(), "$lte": d2.isoformat()}}, {"_id": 0, "photo": 0}).to_list(100000)
        keys = series_keys(d1, d2, group)
        series = {k: {"key": k, "label": bucket_label(k, group), "loans": 0, "returns": 0, "visits": 0} for k in keys}
        books, members = {}, {}

        def member_row(mid, name, cls):
            return members.setdefault(mid, {"member_id": mid, "name": name, "class_name": cls or "", "loans": 0, "visits": 0})

        for l in loans:
            bd = wib_day(l.get("borrowed_at"))
            if bd and d1 <= bd <= d2:
                series[bucket(bd, group)]["loans"] += 1
                b = books.setdefault(l["book_id"], {"book_id": l["book_id"], "title": l.get("book_title", ""), "count": 0})
                b["count"] += 1
                member_row(l["member_id"], l.get("member_name", ""), "")["loans"] += 1
            rd = wib_day(l.get("returned_at"))
            if rd and d1 <= rd <= d2:
                series[bucket(rd, group)]["returns"] += 1
        for v in visits:
            k = bucket(date.fromisoformat(v["date"]), group)
            if k in series:
                series[k]["visits"] += 1
            if v.get("kind") == "anggota" and v.get("member_id"):
                m = member_row(v["member_id"], v.get("name", ""), v.get("class_name", ""))
                m["visits"] += 1
                m["class_name"] = m["class_name"] or v.get("class_name", "")

        info = {m["id"]: m for m in await db.members.find({"id": {"$in": list(members)}}, {"_id": 0, "id": 1, "name": 1, "class_name": 1}).to_list(10000)}
        for mid, m in members.items():
            if mid in info:
                m["name"], m["class_name"] = info[mid]["name"], info[mid].get("class_name", "")
            m["score"] = m["loans"] + m["visits"]

        agg = await db.books.aggregate([{"$group": {"_id": None, "stock": {"$sum": "$stock"}, "titles": {"$sum": 1}}}]).to_list(1)
        a = agg[0] if agg else {"stock": 0, "titles": 0}
        kpis = {
            "books": a["stock"], "titles": a["titles"],
            "members": await db.members.count_documents({}),
            "active_loans": sum(1 for l in loans if l["status"] == "Dipinjam"),
            "overdue": len(await overdue_loans()),
            "visits_today": await db.visits.count_documents({"date": today.isoformat()}),
        }
        rows = list(series.values())
        return {
            "start": d1.isoformat(), "end": d2.isoformat(), "group": group, "kpis": kpis,
            "totals": {"loans": sum(x["loans"] for x in rows), "returns": sum(x["returns"] for x in rows), "visits": sum(x["visits"] for x in rows)},
            "series": rows,
            "top_books": sorted(books.values(), key=lambda b: -b["count"])[:10],
            "top_members": sorted(members.values(), key=lambda m: (-m["score"], -m["loans"]))[:10],
        }

    return r


def _fmt_day(iso):
    d = date.fromisoformat(iso)
    return f"{d.day} {MONTHS[d.month - 1]} {d.year}"


def render_summary_xlsx(d, settings, officer, orientation="landscape"):
    """Laporan Ringkasan & grafik dalam bentuk Excel (.xlsx): ringkasan angka, data per hari/bulan + grafik, buku terpopuler, anggota aktif."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    navy, light = "0F2745", "F1F5F9"
    thin = Side(style="thin", color="CBD5E1"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    now = datetime.now(WIB)
    period = f"{_fmt_day(d['start'])} s.d. {_fmt_day(d['end'])}"
    per = "Per bulan" if d["group"] == "month" else "Per hari"

    def setup(ws, ncol, title, widths):
        for logo_val, fallback, col in ((settings.get("school_logo"), "logo-sekolah.jpg", "A"), (settings.get("library_logo"), "logo-perpustakaan.jpg", get_column_letter(ncol))):
            try:
                img = XLImage(_logo_bytes(logo_val, fallback)); img.width = img.height = 56; ws.add_image(img, f"{col}1")
            except Exception:
                pass
        if ncol - 1 > 2:
            for row in (1, 2, 3):
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=ncol - 1)
        ws["B1"] = "PERPUSTAKAAN SMA NEGERI 1 ABUNG TENGAH"; ws["B1"].font = Font(bold=True, size=14, color=navy)
        ws["B2"] = "Jln Raya Subik, Desa Pekurun Selatan, Lampung Utara"; ws["B2"].font = Font(size=9, color="475569")
        ws["B3"] = f"{title} — {period} · dicetak {fmt_long(now, True)} WIB oleh {officer}"; ws["B3"].font = Font(size=9, italic=True, color="475569")
        for c in ("B1", "B2", "B3"):
            ws[c].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 26; ws.row_dimensions[2].height = 15; ws.row_dimensions[3].height = 28
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.page_setup.orientation = orientation; ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = ws.page_margins.right = 0.4; ws.page_margins.top = ws.page_margins.bottom = 0.5
        ws.print_options.horizontalCentered = True

    def table(ws, row, headers, rows, title=None):
        if title:
            ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=11, color=navy); row += 1
        for j, h in enumerate(headers, 1):
            c = ws.cell(row=row, column=j, value=h); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=navy)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = border
        head = row
        for i, r in enumerate(rows, 1):
            for j, v in enumerate(r, 1):
                c = ws.cell(row=row + i, column=j, value=v); c.border = border; c.alignment = Alignment(vertical="center", wrap_text=True)
                if i % 2 == 0: c.fill = PatternFill("solid", fgColor=light)
        if not rows:
            ws.cell(row=row + 1, column=1, value="Belum ada data pada rentang ini").font = Font(italic=True, color="94A3B8")
        return head, row + max(len(rows), 1)

    wb = Workbook()
    # 1. Ringkasan
    ws = wb.active; ws.title = "Ringkasan"
    setup(ws, 4, "Laporan Ringkasan & Grafik", [26, 18, 26, 18])
    k, t = d["kpis"], d["totals"]
    _, end = table(ws, 5, ["Keterangan", "Jumlah", "Keterangan", "Jumlah"], [
        ["Total buku (eksemplar)", k["books"], "Peminjaman pada rentang", t["loans"]],
        ["Judul buku", k["titles"], "Pengembalian pada rentang", t["returns"]],
        ["Anggota terdaftar", k["members"], "Pengunjung pada rentang", t["visits"]],
        ["Peminjaman aktif", k["active_loans"], "Buku terlambat", k["overdue"]],
        ["Pengunjung hari ini", k["visits_today"], "Pengelompokan", per],
    ], "Ringkasan angka")
    ws.cell(row=end + 2, column=1, value="Grafik tersedia di lembar \"Grafik\"; data rinci di lembar Buku terpopuler & Anggota aktif.").font = Font(italic=True, size=9, color="475569")

    # 2. Grafik (data per hari / bulan + grafik batang & garis)
    gs = wb.create_sheet("Grafik")
    setup(gs, 4, f"Sirkulasi & Pengunjung {per.lower()}", [18, 14, 14, 14])
    series = d["series"]
    head, last = table(gs, 5, ["Periode", "Peminjaman", "Pengembalian", "Pengunjung"], [[x["label"], x["loans"], x["returns"], x["visits"]] for x in series], f"Data {per.lower()}")
    gs.freeze_panes = gs.cell(row=head + 1, column=1)
    if series:
        cats = Reference(gs, min_col=1, min_row=head + 1, max_row=last)
        bar = BarChart(); bar.type = "col"; bar.title = "Peminjaman & pengembalian"; bar.y_axis.title = "Jumlah"; bar.height = 8; bar.width = 17 if orientation == "landscape" else 13
        bar.add_data(Reference(gs, min_col=2, max_col=3, min_row=head, max_row=last), titles_from_data=True); bar.set_categories(cats)
        bar.series[0].graphicalProperties.solidFill = "7C3AED"; bar.series[1].graphicalProperties.solidFill = "16A34A"
        line = LineChart(); line.title = "Tren pengunjung"; line.y_axis.title = "Pengunjung"; line.height = 8; line.width = bar.width
        line.add_data(Reference(gs, min_col=4, max_col=4, min_row=head, max_row=last), titles_from_data=True); line.set_categories(cats)
        line.series[0].graphicalProperties.line.solidFill = "D97706"; line.series[0].graphicalProperties.line.width = 28000
        gs.add_chart(bar, "F5"); gs.add_chart(line, "F23")

    # 3. Buku terpopuler
    bs = wb.create_sheet("Buku terpopuler")
    setup(bs, 3, "Buku terpopuler", [12, 64, 14])
    table(bs, 5, ["Peringkat", "Judul", "Dipinjam"], [[i + 1, b["title"], b["count"]] for i, b in enumerate(d["top_books"])], "Paling sering dipinjam")

    # 4. Anggota aktif
    ms = wb.create_sheet("Anggota aktif")
    setup(ms, 5, "Anggota paling aktif", [12, 34, 16, 14, 14])
    table(ms, 5, ["Peringkat", "Nama", "Kelas", "Peminjaman", "Pengunjung"], [[i + 1, m["name"], m.get("class_name") or "-", m["loans"], m["visits"]] for i, m in enumerate(d["top_members"])], "Peminjaman + pengunjung")

    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def render_summary_pdf(d, settings, officer, orientation="portrait"):
    """Laporan Ringkasan & grafik dalam PDF: kop, kartu angka, grafik sirkulasi & pengunjung, buku terpopuler, anggota aktif, tanda tangan."""
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.charts.legends import Legend
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab_helpers import NAVY, BLUE, LIGHT, logo_flowable, fit_widths, base_table_style

    portrait = orientation == "portrait"
    pagesize = A4 if portrait else landscape(A4)
    mx = 12 * mm if portrait else 14 * mm
    avail = pagesize[0] - 2 * mx
    now = datetime.now(WIB)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=pagesize, leftMargin=mx, rightMargin=mx, topMargin=12 * mm, bottomMargin=14 * mm, title="Laporan Ringkasan & Grafik")
    fs = 8 if portrait else 8.5
    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=13.5 if portrait else 15, textColor=NAVY, leading=17 if portrait else 18)
    h2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=10.5, textColor=NAVY, leading=13, spaceBefore=4, spaceAfter=4)
    sub = ParagraphStyle("sub", fontName="Helvetica", fontSize=8.5, textColor=colors.HexColor("#475569"), leading=11.5)
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=fs, leading=fs + 2.5)
    cellb = ParagraphStyle("cellb", parent=cell, fontName="Helvetica-Bold")
    head = ParagraphStyle("head", fontName="Helvetica-Bold", fontSize=fs - 0.5, leading=fs + 1.5, textColor=colors.white)
    kl = ParagraphStyle("kl", fontName="Helvetica", fontSize=7.5, textColor=colors.HexColor("#64748b"), leading=9)
    kv = ParagraphStyle("kv", fontName="Helvetica-Bold", fontSize=15 if portrait else 16, textColor=NAVY, leading=18)

    logo = 15 * mm
    kop = Table([[logo_flowable(settings.get("school_logo"), "logo-sekolah.jpg", logo), [Paragraph("PERPUSTAKAAN SMA NEGERI 1 ABUNG TENGAH", h1), Paragraph("Jln Raya Subik, Desa Pekurun Selatan, Lampung Utara", sub)], logo_flowable(settings.get("library_logo"), "logo-perpustakaan.jpg", logo)]], colWidths=[logo + 4 * mm, avail - 2 * (logo + 4 * mm), logo + 4 * mm])
    kop.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 0), (-1, 0), 1.5, NAVY), ("BOTTOMPADDING", (0, 0), (-1, 0), 6)]))
    per = "per bulan" if d["group"] == "month" else "per hari"
    title = Paragraph("Laporan Ringkasan &amp; Grafik", ParagraphStyle("t", parent=h1, fontSize=12 if portrait else 13, spaceBefore=8))
    meta = Paragraph(f"Periode {_fmt_day(d['start'])} s.d. {_fmt_day(d['end'])} · data {per} · dicetak {fmt_long(now, True)} WIB oleh {officer}", sub)

    k, t = d["kpis"], d["totals"]
    cards = [("Total buku", k["books"], f"{k['titles']} judul"), ("Anggota", k["members"], "terdaftar"), ("Peminjaman aktif", k["active_loans"], "belum dikembalikan"), ("Buku terlambat", k["overdue"], "lewat batas pinjam"), ("Pengunjung hari ini", k["visits_today"], "tercatat"),
             ("Peminjaman", t["loans"], "pada rentang ini"), ("Pengembalian", t["returns"], "pada rentang ini"), ("Pengunjung", t["visits"], "pada rentang ini")]
    per_row = 4 if portrait else 8
    cells = [[Paragraph(lbl, kl), Paragraph(f"{val:,}".replace(",", "."), kv), Paragraph(note, kl)] for lbl, val, note in cards]
    rows = [cells[i:i + per_row] for i in range(0, len(cells), per_row)]
    rows = [r + [""] * (per_row - len(r)) for r in rows]
    kpi = Table(rows, colWidths=[avail / per_row] * per_row)
    kpi.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")), ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6), ("BACKGROUND", (0, 0), (-1, -1), colors.white)]))

    series = d["series"]
    labels = [x["label"] for x in series]
    step = max(1, -(-len(labels) // (10 if portrait else 8)))  # pembulatan ke atas: label sumbu tidak berdesakan
    show = [l if i % step == 0 else "" for i, l in enumerate(labels)]
    chart_w = avail if portrait else (avail - 6 * mm) / 2
    chart_h = 62 * mm

    def drawing(title_txt):
        dr = Drawing(chart_w, chart_h)
        dr.add(String(0, chart_h - 10, title_txt, fontName="Helvetica-Bold", fontSize=9, fillColor=NAVY))
        return dr

    def axis(c, values):
        top = max([1] + values)
        c.valueAxis.valueMin = 0; c.valueAxis.valueMax = top + max(1, round(top * 0.15))
        c.valueAxis.valueStep = max(1, round((top + 1) / 5))
        c.valueAxis.labels.fontSize = 7; c.valueAxis.labels.fontName = "Helvetica"
        c.valueAxis.gridStrokeColor = colors.HexColor("#e2e8f0"); c.valueAxis.visibleGrid = True
        c.categoryAxis.categoryNames = show or [""]; c.categoryAxis.labels.fontSize = 6.5; c.categoryAxis.labels.fontName = "Helvetica"; c.categoryAxis.labels.angle = 0 if len(show) <= 8 else 35
        c.categoryAxis.labels.boxAnchor = "n" if len(show) <= 8 else "ne"

    bar_d = drawing(f"Peminjaman & pengembalian {per}")
    bc = VerticalBarChart(); bc.x, bc.y = 26, 30; bc.width, bc.height = chart_w - 34, chart_h - 58
    loans_v = [x["loans"] for x in series] or [0]; ret_v = [x["returns"] for x in series] or [0]
    bc.data = [loans_v, ret_v]; axis(bc, loans_v + ret_v)
    bc.bars[0].fillColor = colors.HexColor("#7c3aed"); bc.bars[1].fillColor = colors.HexColor("#16a34a"); bc.bars.strokeColor = None
    bc.groupSpacing = 3 if len(series) > 20 else 6; bc.barSpacing = 0.5
    bar_d.add(bc)
    lg = Legend(); lg.x, lg.y = 30, chart_h - 18; lg.alignment = "right"; lg.columnMaximum = 1; lg.fontSize = 7.5; lg.fontName = "Helvetica"; lg.deltax = 70; lg.boxAnchor = "nw"
    lg.colorNamePairs = [(colors.HexColor("#7c3aed"), "Peminjaman"), (colors.HexColor("#16a34a"), "Pengembalian")]; lg.dxTextSpace = 4; lg.dx = 7; lg.dy = 7
    bar_d.add(lg)

    line_d = drawing(f"Tren pengunjung {per}")
    lc = HorizontalLineChart(); lc.x, lc.y = 26, 30; lc.width, lc.height = chart_w - 34, chart_h - 58
    vis_v = [x["visits"] for x in series] or [0]
    lc.data = [vis_v]; axis(lc, vis_v)
    lc.lines[0].strokeColor = colors.HexColor("#d97706"); lc.lines[0].strokeWidth = 1.8
    line_d.add(lc)

    charts = Table([[bar_d], [line_d]] if portrait else [[bar_d, line_d]], colWidths=[chart_w] if portrait else [chart_w + 3 * mm, chart_w + 3 * mm])
    charts.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")), ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")), ("TOPPADDING", (0, 0), (-1, -1), 6), ("LEFTPADDING", (0, 0), (-1, -1), 4)]))

    def ranked(headers, data, widths):
        body = [[Paragraph(h, head) for h in headers]] + [[Paragraph(str(v), cellb if j == 1 else cell) for j, v in enumerate(r)] for r in data]
        if not data:
            body.append([Paragraph("Belum ada data pada rentang ini", cell)] + [""] * (len(headers) - 1))
        tb = Table(body, colWidths=fit_widths(widths, avail if portrait else (avail - 6 * mm) / 2), repeatRows=1)
        tb.setStyle(base_table_style()); return tb

    books = ranked(["#", "Judul", "Dipinjam"], [[i + 1, b["title"], f"{b['count']}×"] for i, b in enumerate(d["top_books"])], [8, 82, 18])
    members = ranked(["#", "Nama", "Kelas", "Pinjam", "Kunjungan"], [[i + 1, m["name"], m.get("class_name") or "-", m["loans"], m["visits"]] for i, m in enumerate(d["top_members"])], [8, 44, 22, 16, 18])
    tops = [KeepTogether([Paragraph("Buku terpopuler", h2), books]), Spacer(1, 8), KeepTogether([Paragraph("Anggota paling aktif", h2), members])] if portrait else [
        Table([[[Paragraph("Buku terpopuler", h2), books], [Paragraph("Anggota paling aktif", h2), members]]], colWidths=[(avail) / 2] * 2, style=TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 3 * mm)]))]

    sign = Table([["", Paragraph(f"Pekurun Selatan, {fmt_long(now)}<br/>Petugas Perpustakaan,<br/><br/><br/><br/><b>{officer}</b>", ParagraphStyle("s", parent=cell, fontSize=9, leading=12, alignment=1))]], colWidths=[avail - 70 * mm, 70 * mm])
    doc.build([kop, title, meta, Spacer(1, 8), kpi, Spacer(1, 10), KeepTogether([Paragraph("Grafik", h2), charts]), Spacer(1, 10), *tops, Spacer(1, 14), KeepTogether([sign])])
    return buf.getvalue()
