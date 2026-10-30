"""Iteration 6 backend tests: overdue, xlsx export, reservations, stocktake."""
import os, io, uuid, pytest, requests
from openpyxl import load_workbook
from pathlib import Path

def _load_env():
    for p in [Path("/app/frontend/.env"), Path("/app/backend/.env")]:
        if p.exists():
            for line in p.read_text().splitlines():
                if "=" in line and not line.strip().startswith("#"):
                    k, v = line.split("=", 1); os.environ.setdefault(k.strip(), v.strip())
_load_env()
BASE = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")


def login(portal, email, password):
    s = requests.Session()
    r = s.post(f"{BASE}/api/auth/login", json={"email": email, "password": password, "portal": portal}, timeout=15)
    assert r.status_code == 200, r.text
    return s


@pytest.fixture(scope="module")
def petugas():
    return login("admin", "petugas@smansa.sch.id", "petugas123")


@pytest.fixture(scope="module")
def siswa():
    return login("siswa", "siswa@smansa.sch.id", "siswa123")


# ---------- overdue + dashboard ----------
class TestOverdue:
    def test_overdue_staff(self, petugas):
        r = petugas.get(f"{BASE}/api/overdue")
        assert r.status_code == 200
        d = r.json()
        assert d["days"] == 3
        bk1043 = [i for i in d["items"] if i["book_id"] == "BK-1043"]
        assert bk1043, "BK-1043 overdue loan expected"
        assert bk1043[0]["days_out"] >= 10

    def test_overdue_siswa_forbidden(self, siswa):
        assert siswa.get(f"{BASE}/api/overdue").status_code == 403

    def test_dashboard(self, petugas):
        d = petugas.get(f"{BASE}/api/dashboard").json()
        assert d["overdue"] >= 1
        assert isinstance(d["reservations_waiting"], int)


# ---------- xlsx export ----------
class TestXlsxExport:
    @pytest.mark.parametrize("kind", ["transaksi", "masuk", "hilang"])
    def test_export(self, petugas, kind):
        r = petugas.get(f"{BASE}/api/reports/export.xlsx", params={"type": kind, "period": "semua"})
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        assert "attachment" in r.headers.get("content-disposition", "")
        assert ".xlsx" in r.headers.get("content-disposition", "")
        assert r.content[:2] == b"PK"
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        headers = [ws.cell(row=7, column=c).value for c in range(1, 10) if ws.cell(row=7, column=c).value]
        if kind == "transaksi":
            assert headers == ["No", "Anggota", "Kelas", "Buku", "Tgl Pinjam", "Tgl Kembali", "Status"]

    def test_invalid_type(self, petugas):
        assert petugas.get(f"{BASE}/api/reports/export.xlsx", params={"type": "foo"}).status_code == 400

    def test_siswa_forbidden(self, siswa):
        assert siswa.get(f"{BASE}/api/reports/export.xlsx").status_code == 403


# ---------- Reservations ----------
def _own_book(petugas):
    # stock-1 book owned by this test so parallel workers cannot touch it
    return petugas.post(f"{BASE}/api/books", json={"title": f"RSV_{uuid.uuid4().hex[:6]}", "author": "T", "shelf": "R-01"}).json()["id"]


class TestReservationsSiswa:
    @classmethod
    def _cleanup(cls, petugas):
        # cancel all Menunggu reservations for member Bima
        rows = petugas.get(f"{BASE}/api/reservations").json()
        for r in rows:
            if r["member_id"] == "MB-2024001" and r["status"] == "Menunggu":
                petugas.delete(f"{BASE}/api/reservations/{r['id']}")

    def test_full_flow(self, siswa, petugas):
        self._cleanup(petugas)
        bk = _own_book(petugas)

        r = siswa.post(f"{BASE}/api/reservations", json={"book_id": bk})
        assert r.status_code in (200, 201), r.text
        res = r.json()
        assert res["status"] == "Menunggu"
        assert res["member_id"] == "MB-2024001"
        rid = res["id"]

        # duplicate
        r2 = siswa.post(f"{BASE}/api/reservations", json={"book_id": bk})
        assert r2.status_code == 400

        # mine includes
        mine = siswa.get(f"{BASE}/api/reservations/mine").json()
        assert any(x["id"] == rid for x in mine)

        # /api/books waiting=1
        books = siswa.get(f"{BASE}/api/books").json()
        row = [b for b in books if b["id"] == bk][0]
        assert row["waiting"] >= 1

        # cancel
        assert siswa.delete(f"{BASE}/api/reservations/{rid}").status_code == 200
        # cancel again -> 400
        assert siswa.delete(f"{BASE}/api/reservations/{rid}").status_code == 400

    def test_hilang_book_cannot_reserve(self, siswa, petugas):
        # mark a book Hilang temporarily
        bk = _own_book(petugas)
        petugas.patch(f"{BASE}/api/books/{bk}/status", json={"status": "Hilang", "note": "test", "reporter": "T"})
        try:
            r = siswa.post(f"{BASE}/api/reservations", json={"book_id": bk})
            assert r.status_code == 400
        finally:
            petugas.patch(f"{BASE}/api/books/{bk}/status", json={"status": "Tersedia"})

    def test_max_3(self, siswa, petugas):
        self._cleanup(petugas)
        # own books so parallel tests cannot delete them mid-test
        avail = [_own_book(petugas) for _ in range(4)]
        created = []
        try:
            for b in avail[:3]:
                r = siswa.post(f"{BASE}/api/reservations", json={"book_id": b})
                assert r.status_code in (200, 201), r.text
                created.append(r.json()["id"])
            r4 = siswa.post(f"{BASE}/api/reservations", json={"book_id": avail[3]})
            assert r4.status_code == 400
            assert "Maksimal 3" in r4.json().get("detail", "")
        finally:
            for rid in created:
                siswa.delete(f"{BASE}/api/reservations/{rid}")
            for b in avail:
                petugas.delete(f"{BASE}/api/books/{b}")

    # Same class as test_max_3 so xdist loadscope keeps reservation tests on one worker.
    def test_process_pinjam_and_tolak(self, siswa, petugas):
        # cleanup
        TestReservationsSiswa._cleanup(petugas)
        bk_ok = _own_book(petugas)
        books = petugas.get(f"{BASE}/api/books").json()

        # siswa creates reservation
        r1 = siswa.post(f"{BASE}/api/reservations", json={"book_id": bk_ok}).json()
        rid1 = r1["id"]

        # staff lists
        lst = petugas.get(f"{BASE}/api/reservations").json()
        assert lst[0]["status"] == "Menunggu"

        # process pinjam
        # one title = one record with stock; borrowing lowers available count
        avail_before = [b for b in books if b["id"] == bk_ok][0]["available"]
        r = petugas.post(f"{BASE}/api/reservations/{rid1}/process", json={"action": "pinjam"})
        assert r.status_code == 200
        assert r.json()["status"] == "Dipinjamkan"
        books = petugas.get(f"{BASE}/api/books").json()
        assert [b for b in books if b["id"] == bk_ok][0]["available"] == avail_before - 1
        # new transaction exists
        tx = petugas.get(f"{BASE}/api/transactions").json()
        assert any(t["book_id"] == bk_ok and t["action"] == "pinjam" and t.get("reservation_id") == rid1 for t in tx)

        # process again -> 404 (already processed)
        r_again = petugas.post(f"{BASE}/api/reservations/{rid1}/process", json={"action": "pinjam"})
        assert r_again.status_code == 404

        # reservation on Dipinjam book -> pinjam process returns 400
        r2 = siswa.post(f"{BASE}/api/reservations", json={"book_id": bk_ok}).json()
        rid2 = r2["id"]
        rp = petugas.post(f"{BASE}/api/reservations/{rid2}/process", json={"action": "pinjam"})
        assert rp.status_code == 400

        # invalid action
        inv = petugas.post(f"{BASE}/api/reservations/{rid2}/process", json={"action": "foo"})
        assert inv.status_code == 400

        # tolak
        tolak = petugas.post(f"{BASE}/api/reservations/{rid2}/process", json={"action": "tolak"})
        assert tolak.status_code == 200
        assert tolak.json()["status"] == "Ditolak"

        # return book to Tersedia (manual kembali tx)
        petugas.post(f"{BASE}/api/transactions", json={"member_id": "MB-2024001", "book_id": bk_ok, "action": "kembali"})

    def test_siswa_forbidden(self, siswa):
        assert siswa.get(f"{BASE}/api/reservations").status_code == 403
        assert siswa.post(f"{BASE}/api/reservations/RS-FAKE/process", json={"action": "pinjam"}).status_code == 403


# ---------- Stocktake ----------
class TestStocktake:
    def test_save_and_list(self, petugas):
        r = petugas.post(f"{BASE}/api/stocktake", json={"shelf": "R-03", "scanned": ["BK-1045"], "unscanned": ["BK-1042"], "misplaced": []})
        assert r.status_code == 200
        d = r.json()
        assert d["id"].startswith("SO-")
        assert d["scanned_count"] == 1
        assert d["unscanned_count"] == 1
        lst = petugas.get(f"{BASE}/api/stocktake").json()
        assert lst[0]["id"] == d["id"]

    def test_siswa_forbidden(self, siswa):
        assert siswa.get(f"{BASE}/api/stocktake").status_code == 403
        assert siswa.post(f"{BASE}/api/stocktake", json={"shelf": "R-01"}).status_code == 403


# ---------- Regression /siswa isolation ----------
class TestSiswaIsolation:
    @pytest.mark.parametrize("path", ["/api/members", "/api/reports", "/api/transactions"])
    def test_forbidden(self, siswa, path):
        assert siswa.get(f"{BASE}{path}").status_code == 403
