import os
import uuid

import pytest
import requests
from openpyxl import load_workbook
import io

BASE = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")


def login(email, password, portal):
    s = requests.Session()
    r = s.post(f"{BASE}/api/auth/login", json={"email": email, "password": password, "portal": portal})
    assert r.status_code == 200, r.text
    s.headers["Authorization"] = f"Bearer {r.json()['token']}"
    return s


@pytest.fixture(scope="module")
def petugas():
    return login("petugas@smansa.sch.id", "petugas123", "admin")


@pytest.fixture(scope="module")
def siswa():
    return login("siswa@smansa.sch.id", "siswa123", "siswa")


class TestLoanLeaderboard:
    def test_ranking_counts_loans_per_student_and_class(self, petugas):
        tag = uuid.uuid4().hex[:6]
        cls = f"TEST_LB_{tag}"
        a = petugas.post(f"{BASE}/api/members", json={"name": f"TEST_A_{tag}", "student_id": f"A{tag}", "class_name": cls}).json()["id"]
        b = petugas.post(f"{BASE}/api/members", json={"name": f"TEST_B_{tag}", "student_id": f"B{tag}", "class_name": cls}).json()["id"]
        books = [petugas.post(f"{BASE}/api/books", json={"title": f"TEST_LB_{tag}_{i}", "author": "T", "shelf": "R-01"}).json()["id"] for i in range(3)]
        for mid, bid in ((a, books[0]), (a, books[1]), (b, books[2])):
            r = petugas.post(f"{BASE}/api/transactions", json={"member_id": mid, "book_id": bid, "action": "pinjam"})
            assert r.status_code == 200, r.text
        lb = petugas.get(f"{BASE}/api/reports/loan-leaderboard").json()
        mine = {s["member_id"]: s for s in lb["students"] if s["member_id"] in (a, b)}
        assert mine[a]["loans"] == 2 and mine[a]["titles"] == 2
        assert mine[b]["loans"] == 1 and mine[a]["rank"] < mine[b]["rank"]
        c = next(c for c in lb["classes"] if c["class_name"] == cls)
        assert c["loans"] == 3 and c["members"] == 2
        assert lb["total_loans"] >= 3 and lb["month"] in lb["months"]

    def test_invalid_month_rejected(self, petugas):
        assert petugas.get(f"{BASE}/api/reports/loan-leaderboard", params={"month": "2026-13"}).status_code == 400

    def test_siswa_forbidden(self, siswa):
        assert siswa.get(f"{BASE}/api/reports/loan-leaderboard").status_code == 403

    def test_export_xlsx_and_pdf(self, petugas):
        month = petugas.get(f"{BASE}/api/reports/loan-leaderboard").json()["month"]
        x = petugas.get(f"{BASE}/api/reports/export.xlsx", params={"type": "peminjam", "period": month})
        assert x.status_code == 200
        ws = load_workbook(io.BytesIO(x.content)).active
        assert [c.value for c in ws[7]] == ["Peringkat", "Nama Siswa", "Kelas", "Peminjaman", "Judul berbeda", "Persentase"]
        p = petugas.get(f"{BASE}/api/reports/export.pdf", params={"type": "peminjam", "period": month})
        assert p.status_code == 200 and p.content[:4] == b"%PDF"
