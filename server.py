from dotenv import load_dotenv
load_dotenv(override=False)

from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, Response, UploadFile, File, Form
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field
from typing import Optional, List
from datetime import datetime, timezone, timedelta
from pathlib import Path
import os, re, uuid, bcrypt, jwt, io, csv, json, logging, secrets
from openpyxl import load_workbook
from media import init_storage, photo_upload_allowed, build_router as media_router
from reports_pdf import WIB, build_loans, in_period, render_report_pdf, render_report_xlsx, fmt_long

ROOT_DIR = Path(__file__).parent
mongo_url = os.environ["MONGO_URL"]
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ["DB_NAME"]]
app = FastAPI(title="Perpustakaan SMA Negeri 1 Abung Tengah")
api = APIRouter(prefix="/api")
JWT_SECRET = os.environ["JWT_SECRET"]
ADMIN_NAME = "Devi Latifa, S.I.Pust"
if len(JWT_SECRET) < 32:
    raise RuntimeError("JWT_SECRET must contain at least 32 characters")

class LoginInput(BaseModel):
    email: str
    password: str
    portal: str = "admin"

class BookInput(BaseModel):
    title: str
    author: str
    category: str = "Umum"
    shelf: str = "R-01"
    isbn: str = ""
    cover: str = ""
    stock: int = Field(1, ge=1, le=1000)

class BookUpdate(BaseModel):
    title: str = Field(..., min_length=1, max_length=200)
    author: str = Field(..., min_length=1, max_length=120)
    category: str = Field("Umum", max_length=60)
    shelf: str = Field(..., min_length=1, max_length=20)
    isbn: str = Field("", max_length=40)
    cover: str = ""

class ShelfInput(BaseModel):
    id: str = Field(..., pattern=r"^[A-Za-z0-9-]{1,12}$")
    name: str = Field(..., min_length=1, max_length=60)
    capacity: int = Field(..., ge=1, le=10000)
    color: str = Field("azure", max_length=20)

class StockInput(BaseModel):
    stock: int = Field(..., ge=0, le=5000)

class MemberInput(BaseModel):
    name: str
    student_id: str
    class_name: str
    photo: str = ""

class BulkDeleteInput(BaseModel):
    class_name: str

class PromoteInput(BaseModel):
    from_class: str
    to_class: str

class PhotoInput(BaseModel):
    photo: str = ""

class ResetInput(BaseModel):
    include_books: bool = False
    code: str = ""

class PasswordChangeInput(BaseModel):
    current_password: str
    new_password: str = Field(..., min_length=6)

class PasswordSetInput(BaseModel):
    new_password: str = Field(..., min_length=6)

class TransactionInput(BaseModel):
    member_id: str
    book_id: str
    action: str = "pinjam"

class IncomingInput(BaseModel):
    title: str
    author: str
    category: str = "Umum"
    shelf: str = "R-01"
    isbn: str = ""
    cover: str = ""
    quantity: int = Field(1, ge=1, le=200)
    date: str = ""
    source: str = ""

class StatusInput(BaseModel):
    status: str
    note: str = ""
    reporter: str = ""

class SettingsInput(BaseModel):
    school_logo: str = ""
    library_logo: str = ""
    card_theme: str = "navy"
    card_layout: str = "classic"
    card_show_qr: bool = True
    card_show_address: bool = True
    card_title: str = "KARTU ANGGOTA PERPUSTAKAAN"
    card_school_name: str = "SMA NEGERI 1 ABUNG TENGAH"
    card_font: str = "modern"
    card_font_size: str = "sedang"
    card_school_spacing: float = 0
    card_school_logo_pos: str = "kiri"
    card_library_logo_pos: str = "kanan"
    card_photo_pos: str = "kiri"
    card_qr_pos: str = "atas"
    card_back_title: str = Field("PERATURAN & TATA TERTIB PERPUSTAKAAN", max_length=120)
    card_back_rules: str = Field("Kartu ini wajib dibawa setiap meminjam dan mengembalikan buku.\nKembalikan buku tepat waktu sesuai tanggal yang ditentukan petugas.\nJaga buku tetap bersih; dilarang melipat, mencoret, atau merobek.\nBuku yang hilang atau rusak wajib diganti dengan judul yang sama.\nKartu tidak boleh dipinjamkan kepada orang lain.\nJika kartu ini ditemukan, mohon kembalikan ke Perpustakaan SMANSA.", max_length=2000)
    card_back_web: str = Field("perpustakaan.sman1abungtengah.sch.id", max_length=150)

DEFAULT_SETTINGS = SettingsInput().model_dump()

def staff_only(user):
    if user["role"] not in ["admin", "petugas"]: raise HTTPException(403, "Akses ditolak")

def now_iso():
    return datetime.now(timezone.utc).isoformat()

def public(doc):
    doc = dict(doc)
    doc.pop("_id", None)
    doc.pop("password_hash", None)
    return doc

def password_hash(value):
    return bcrypt.hashpw(value.encode(), bcrypt.gensalt()).decode()

def book_key(b):
    return (str(b.get("title", "")).strip().lower(), str(b.get("author", "")).strip().lower())

def compute_state(b):
    stock = max(int(b.get("stock") or 0), 0); borrowed = max(int(b.get("borrowed") or 0), 0); lost = max(int(b.get("lost") or 0), 0)
    borrowed = min(borrowed, stock)
    available = stock - borrowed
    status = "Tersedia" if available > 0 else ("Dipinjam" if borrowed > 0 else ("Hilang" if lost > 0 else "Habis"))
    return {"stock": stock, "borrowed": borrowed, "lost": lost, "available": available, "status": status}

async def refresh_book(book_id):
    doc = await db.books.find_one({"id": book_id}, {"_id": 0})
    if not doc: return None
    state = compute_state(doc)
    await db.books.update_one({"id": book_id}, {"$set": state})
    return {**doc, **state}

async def find_same_book(title, author):
    key = (title.strip().lower(), author.strip().lower())
    rx = {"$regex": f"^\\s*{re.escape(title.strip())}\\s*$", "$options": "i"}
    for b in await db.books.find({"title": rx}, {"_id": 0}).to_list(50):
        if book_key(b) == key: return b
    return None

async def normalize_books():
    """Gabungkan buku berjudul & penulis sama menjadi satu judul dengan angka stok (idempoten)."""
    books = await db.books.find({}, {"_id": 0}).sort("created_at", 1).to_list(100000)
    groups = {}
    for b in books: groups.setdefault(book_key(b), []).append(b)
    for items in groups.values():
        if len(items) == 1 and "stock" in items[0] and "available" in items[0]: continue
        main = items[0]; stock = borrowed = lost = 0
        for b in items:
            if "stock" in b:
                stock += int(b.get("stock") or 0); borrowed += int(b.get("borrowed") or 0); lost += int(b.get("lost") or 0)
            else:
                st = b.get("status", "Tersedia")
                if st == "Hilang": lost += 1
                else:
                    stock += 1
                    if st == "Dipinjam": borrowed += 1
        extra = {"cover": next((b.get("cover") for b in items if b.get("cover")), ""), "isbn": next((b.get("isbn") for b in items if b.get("isbn")), "")}
        await db.books.update_one({"id": main["id"]}, {"$set": {**compute_state({"stock": stock, "borrowed": borrowed, "lost": lost}), **extra}})
        others = [b["id"] for b in items[1:]]
        if not others: continue
        await db.books.delete_many({"id": {"$in": others}})
        for coll in ("transactions", "reservations", "missing"):
            await db[coll].update_many({"book_id": {"$in": others}}, {"$set": {"book_id": main["id"]}})
        await db.missing.update_many({"code": {"$in": others}}, {"$set": {"code": main["id"]}})
        for inc in await db.incoming.find({"book_ids": {"$in": others}}, {"_id": 0, "id": 1, "book_ids": 1}).to_list(5000):
            ids = list(dict.fromkeys(main["id"] if i in others else i for i in inc.get("book_ids", [])))
            await db.incoming.update_one({"id": inc["id"]}, {"$set": {"book_ids": ids}})

async def open_loan_count(member_id, book_id):
    p = await db.transactions.count_documents({"member_id": member_id, "book_id": book_id, "action": "pinjam"})
    k = await db.transactions.count_documents({"member_id": member_id, "book_id": book_id, "action": "kembali"})
    return p - k

async def take_copy(book_id):
    res = await db.books.find_one_and_update({"id": book_id, "$expr": {"$gt": ["$stock", "$borrowed"]}}, {"$inc": {"borrowed": 1}})
    if not res: raise HTTPException(400, "Stok buku habis — semua eksemplar sedang dipinjam")
    return await refresh_book(book_id)

async def return_copy(book_id):
    res = await db.books.find_one_and_update({"id": book_id, "borrowed": {"$gt": 0}}, {"$inc": {"borrowed": -1}})
    if not res: raise HTTPException(400, "Tidak ada eksemplar buku ini yang sedang dipinjam")
    return await refresh_book(book_id)

def token_for(user):
    return jwt.encode({"sub": user["id"], "exp": datetime.now(timezone.utc) + timedelta(hours=8)}, JWT_SECRET, algorithm="HS256")

async def current_user(request: Request):
    token = request.cookies.get("access_token") or request.headers.get("Authorization", "").replace("Bearer ", "")
    if not token:
        raise HTTPException(401, "Silakan masuk terlebih dahulu")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=["HS256"])
        user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0})
        if not user: raise HTTPException(401, "Sesi tidak ditemukan")
        return user
    except jwt.PyJWTError:
        raise HTTPException(401, "Sesi telah berakhir")

async def seed():
    if await db.users.count_documents({}) == 0:
        users = [
            {"id":"u-admin","name":ADMIN_NAME,"email":"admin@smansa.sch.id","role":"admin","password_hash":password_hash("admin123")},
            {"id":"u-petugas","name":"Pustakawan","email":"petugas@smansa.sch.id","role":"petugas","password_hash":password_hash("petugas123")},
            {"id":"u-siswa","name":"Bima Pratama","email":"siswa@smansa.sch.id","role":"siswa","student_id":"24001","password_hash":password_hash("siswa123")},
        ]
        await db.users.insert_many(users)
    # One-time content seed. Gated by meta flag so a manual data reset stays empty permanently.
    meta = await db.meta.find_one({"id": "seed"}, {"_id": 0})
    if not (meta and meta.get("done")):
        if await db.shelves.count_documents({}) == 0:
            await db.shelves.insert_many([{ "id":"R-01", "name":"Pengetahuan Umum", "color":"azure", "capacity":80 }, {"id":"R-02","name":"Sains & Teknologi","color":"mint","capacity":60},{"id":"R-03","name":"Bahasa & Sastra","color":"amber","capacity":50}])
        if await db.books.count_documents({}) == 0:
            await db.books.insert_many([
                {"id":"BK-1042","title":"Laskar Pelangi","author":"Andrea Hirata","category":"Sastra","shelf":"R-03","isbn":"9789793062792","cover":"https://images.unsplash.com/photo-1544947950-fa07a98d237f?auto=format&fit=crop&w=400&q=70","status":"Tersedia","created_at":now_iso()},
                {"id":"BK-1043","title":"Ensiklopedia Sains Remaja","author":"Tim Edukasi","category":"Sains","shelf":"R-02","isbn":"9786021234567","cover":"https://images.unsplash.com/photo-1532012197267-da84d127e765?auto=format&fit=crop&w=400&q=70","status":"Dipinjam","created_at":now_iso()},
                {"id":"BK-1044","title":"Matematika Dasar SMA","author":"D. Kurniawan","category":"Pelajaran","shelf":"R-01","isbn":"9786027654321","cover":"https://images.unsplash.com/photo-1509228468518-180dd4864904?auto=format&fit=crop&w=400&q=70","status":"Tersedia","created_at":now_iso()},
                {"id":"BK-1045","title":"Bumi Manusia","author":"Pramoedya Ananta Toer","category":"Sastra","shelf":"R-03","isbn":"9789799731234","cover":"https://images.unsplash.com/photo-1476275466078-4007374efbbe?auto=format&fit=crop&w=400&q=70","status":"Tersedia","created_at":now_iso()},
            ])
            # Data contoh langsung memakai format stok (hanya menyentuh buku contoh yang baru dimasukkan)
            for bid, borrowed in (("BK-1042", 0), ("BK-1043", 1), ("BK-1044", 0), ("BK-1045", 0)):
                await db.books.update_one({"id": bid}, {"$set": compute_state({"stock": 1, "borrowed": borrowed, "lost": 0})})
        if await db.members.count_documents({}) == 0:
            await db.members.insert_many([{"id":"MB-2024001","name":"Bima Pratama","student_id":"24001","class_name":"XI IPA 2","photo":""},{"id":"MB-2024002","name":"Siti Rahma","student_id":"24002","class_name":"X IPS 1","photo":""},{"id":"MB-2024003","name":"Raka Wijaya","student_id":"24003","class_name":"XII IPA 1","photo":""}])
        if await db.transactions.count_documents({}) == 0:
            await db.transactions.insert_many([{ "id":"TRX-001","member_id":"MB-2024001","member_name":"Bima Pratama","book_id":"BK-1043","book_title":"Ensiklopedia Sains Remaja","action":"pinjam","date":(datetime.now(timezone.utc)-timedelta(days=10)).isoformat()},{"id":"TRX-002","member_id":"MB-2024002","member_name":"Siti Rahma","book_id":"BK-1042","book_title":"Laskar Pelangi","action":"kembali","date":now_iso()}])
        await db.meta.update_one({"id": "seed"}, {"$set": {"id": "seed", "done": True}}, upsert=True)
    # Migrasi data lama — dijalankan SEKALI saja (tidak menyentuh data produksi di setiap restart)
    if not await db.meta.find_one({"id": "legacy-migrations-v1"}):
        await db.members.update_many({}, {"$unset": {"email": "", "phone": ""}})
        await db.users.update_one({"id": "u-petugas", "name": "Rina Pustakawan"}, {"$set": {"name": "Pustakawan"}})
        for coll, field in (("incoming", "recorded_by"), ("missing", "reported_by"), ("reservations", "processed_by"), ("visits", "recorded_by"), ("visits", "checkout_by"), ("stocktake", "officer")):
            await db[coll].update_many({field: "Rina Pustakawan"}, {"$set": {field: "Pustakawan"}})
        await db.settings.update_many({"card_school_name": "SMA 1 NEGERI ABUNG TENGAH"}, {"$set": {"card_school_name": "SMA NEGERI 1 ABUNG TENGAH"}})
        # (penggabungan buku ganda hanya lewat pemulihan cadangan, tidak otomatis saat start)
        await db.users.update_one({"id":"u-siswa","student_id":{"$exists":False}},{"$set":{"student_id":"24001"}})
        await db.meta.update_one({"id": "legacy-migrations-v1"}, {"$set": {"id": "legacy-migrations-v1", "done": True}}, upsert=True)
    await db.users.update_one({"id": "u-admin", "name": "Administrator"}, {"$set": {"name": "Devi Latifa"}})
    # Gelar admin: "Devi Latifa" → "Devi Latifa, S.I.Pust" (sekali saja, juga di catatan petugas)
    if not await db.meta.find_one({"id": "admin-title-v1"}):
        await db.users.update_one({"id": "u-admin", "name": "Devi Latifa"}, {"$set": {"name": ADMIN_NAME}})
        for coll, field in (("incoming", "recorded_by"), ("missing", "reported_by"), ("reservations", "processed_by"), ("visits", "recorded_by"), ("visits", "checkout_by"), ("stocktake", "officer")):
            await db[coll].update_many({field: "Devi Latifa"}, {"$set": {field: ADMIN_NAME}})
        await db.meta.update_one({"id": "admin-title-v1"}, {"$set": {"id": "admin-title-v1", "done": True}}, upsert=True)

async def ensure_indexes():
    try:
        await db.visits.create_index("check_in"); await db.visits.create_index("date"); await db.visits.create_index("member_id")
        await db.members.create_index("id"); await db.members.create_index("student_id")
        await db.books.create_index("id"); await db.transactions.create_index("member_id"); await db.transactions.create_index("book_id")
    except Exception as e:  # indeks hanya untuk kecepatan, jangan gagalkan start
        logging.getLogger(__name__).warning("Gagal membuat indeks: %s", e)

@app.on_event("startup")
async def startup():
    await seed()
    await db.users.create_index("email", unique=True)
    try:
        init_storage()
    except Exception as e:
        print(f"Storage init failed: {e}")
    await ensure_indexes()

@api.get("/")
async def root(): return {"message":"Perpustakaan SMA Negeri 1 Abung Tengah aktif"}

@api.post("/auth/login")
async def login(data: LoginInput, response: Response):
    identifier = data.email.lower()
    attempt = await db.login_attempts.find_one({"identifier": identifier}, {"_id": 0})
    if attempt and attempt.get("locked_until", "") > now_iso():
        raise HTTPException(429, "Terlalu banyak percobaan. Coba lagi beberapa menit.")
    user = await db.users.find_one({"email": identifier}, {"_id": 0})
    if not user or not bcrypt.checkpw(data.password.encode(), user["password_hash"].encode()):
        failures = (attempt or {}).get("failures", 0) + 1
        payload = {"identifier": identifier, "failures": failures}
        if failures >= 5: payload["locked_until"] = (datetime.now(timezone.utc) + timedelta(minutes=15)).isoformat()
        await db.login_attempts.update_one({"identifier": identifier}, {"$set": payload}, upsert=True)
        raise HTTPException(401, "Email atau kata sandi salah")
    await db.login_attempts.delete_one({"identifier": identifier})
    if data.portal == "siswa" and user["role"] != "siswa":
        raise HTTPException(403, "Akun petugas/admin harus masuk melalui portal /admin")
    if data.portal == "admin" and user["role"] == "siswa":
        raise HTTPException(403, "Akun siswa hanya dapat masuk melalui portal /siswa")
    token = token_for(user)
    response.set_cookie("access_token", token, httponly=True, secure=True, samesite="none", max_age=28800, path="/")
    return {**public(user), "token": token}

@api.get("/auth/me")
async def me(user=Depends(current_user)): return public(user)

@api.post("/auth/logout")
async def logout(response: Response): response.delete_cookie("access_token", path="/", secure=True, samesite="none"); return {"ok":True}

@api.post("/account/password")
async def change_own_password(data: PasswordChangeInput, user=Depends(current_user)):
    staff_only(user)
    if not bcrypt.checkpw(data.current_password.encode(), user["password_hash"].encode()):
        raise HTTPException(400, "Kata sandi lama salah")
    await db.users.update_one({"id": user["id"]}, {"$set": {"password_hash": password_hash(data.new_password)}})
    return {"ok": True}

@api.get("/users")
async def list_users(user=Depends(current_user)):
    if user["role"] != "admin": raise HTTPException(403, "Hanya admin yang dapat mengelola akun")
    rows = await db.users.find({"role": {"$in": ["admin", "petugas"]}}, {"_id":0}).to_list(100)
    return [public(r) for r in rows]

@api.put("/users/{user_id}/password")
async def set_user_password(user_id: str, data: PasswordSetInput, user=Depends(current_user)):
    if user["role"] != "admin": raise HTTPException(403, "Hanya admin yang dapat mengubah kata sandi akun lain")
    target = await db.users.find_one({"id": user_id}, {"_id":0})
    if not target: raise HTTPException(404, "Akun tidak ditemukan")
    if target["role"] not in ["admin", "petugas"]: raise HTTPException(400, "Hanya akun admin/petugas yang dapat diubah di sini")
    await db.users.update_one({"id": user_id}, {"$set": {"password_hash": password_hash(data.new_password)}})
    return {"ok": True}

OVERDUE_DAYS = 3

async def overdue_loans():
    tx = await db.transactions.find({}, {"_id":0}).to_list(5000)
    cutoff = (datetime.now(timezone.utc) - timedelta(days=OVERDUE_DAYS)).isoformat()
    now = datetime.now(timezone.utc)
    out = []
    for l in build_loans(tx):
        if l["status"] == "Dipinjam" and l["borrowed_at"] and l["borrowed_at"] < cutoff:
            days = (now - datetime.fromisoformat(l["borrowed_at"])).days
            out.append({**l, "days_out": days})
    out.sort(key=lambda l: -l["days_out"])
    return out

@api.get("/dashboard")
async def dashboard(user=Depends(current_user)):
    incoming = await db.incoming.aggregate([{"$group":{"_id":None,"total":{"$sum":"$quantity"}}}]).to_list(1)
    overdue = len(await overdue_loans()) if user["role"] != "siswa" else 0
    waiting = await db.reservations.count_documents({"status":"Menunggu"}) if user["role"] != "siswa" else 0
    agg = await db.books.aggregate([{"$group":{"_id":None,"stock":{"$sum":"$stock"},"borrowed":{"$sum":"$borrowed"},"lost":{"$sum":"$lost"},"available":{"$sum":"$available"},"titles":{"$sum":1}}}]).to_list(1)
    a = agg[0] if agg else {"stock":0,"borrowed":0,"lost":0,"available":0,"titles":0}
    vis = {"visitors_today":0,"visitors_inside":0}
    if user["role"] != "siswa":
        await visits_auto_close()
        today = datetime.now(WIB).date().isoformat()
        vis = {"visitors_today": await db.visits.count_documents({"date": today}), "visitors_inside": await db.visits.count_documents({"date": today, "check_out": None})}
    return {**vis, "books":a["stock"],"titles":a["titles"],"available":a["available"],"members":await db.members.count_documents({}),"borrowed":a["borrowed"],"missing":a["lost"],"incoming":incoming[0]["total"] if incoming else 0,"transactions":await db.transactions.count_documents({}),"overdue":overdue,"reservations_waiting":waiting}

@api.get("/overdue")
async def overdue(user=Depends(current_user)):
    staff_only(user)
    return {"days": OVERDUE_DAYS, "items": await overdue_loans()}

def own_member_filter(user):
    if user.get("member_id"): return {"id": user["member_id"]}
    return {"$or": [{"student_id": user.get("student_id", "__none__")}, {"name": {"$regex": f"^{re.escape(user['name'])}$", "$options": "i"}}]}

@api.get("/account/member")
async def my_member(user=Depends(current_user)):
    if user["role"] != "siswa":
        return None
    return await db.members.find_one(own_member_filter(user), {"_id":0})

async def check_student_photo(user, photo, member_id=None):
    if user["role"] != "siswa": return
    if not await photo_upload_allowed(db): raise HTTPException(403, "Unggah foto sedang ditutup oleh petugas perpustakaan")
    if photo and not photo.startswith("/api/files/"): raise HTTPException(400, "Foto harus diunggah melalui tombol unggah")
    if member_id and user.get("member_id") and user["member_id"] != member_id: raise HTTPException(403, "Kamu hanya dapat mengubah fotomu sendiri")

@api.post("/account/photo")
async def set_my_photo(data: PhotoInput, user=Depends(current_user)):
    if user["role"] != "siswa":
        raise HTTPException(403, "Hanya siswa yang dapat mengubah fotonya sendiri")
    await check_student_photo(user, data.photo)
    member = await db.members.find_one(own_member_filter(user), {"_id":0})
    await check_student_photo_allowed(member)
    if member:
        await db.members.update_one({"id": member["id"]}, {"$set": {"photo": data.photo}})
        return await db.members.find_one({"id": member["id"]}, {"_id":0})
    # Belum terdaftar sebagai anggota: buat data anggota otomatis dari akun siswa
    doc = {"id": "MB-" + str(uuid.uuid4())[:8].upper(), "name": user["name"], "student_id": user.get("student_id", ""), "class_name": user.get("class_name", ""), "photo": data.photo}
    await db.members.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}

class ReservationInput(BaseModel):
    book_id: str
    member_id: str = ""

class ProcessInput(BaseModel):
    action: str

@api.get("/reservations")
async def reservations(user=Depends(current_user)):
    staff_only(user)
    rows = await db.reservations.find({}, {"_id":0}).sort("created_at", -1).to_list(500)
    rows.sort(key=lambda r: (r["status"] != "Menunggu", r["created_at"] if r["status"] != "Menunggu" else ""), reverse=False)
    return rows

@api.get("/reservations/mine")
async def my_reservations(member_id: str = "", user=Depends(current_user)):
    if member_id and not user.get("member_id"):
        member = await db.members.find_one({"id": member_id}, {"_id":0})
    else:
        member = await db.members.find_one(own_member_filter(user), {"_id":0})
    if not member: return []
    return await db.reservations.find({"member_id": member["id"]}, {"_id":0}).sort("created_at", -1).to_list(100)

@api.post("/reservations")
async def reserve(data: ReservationInput, user=Depends(current_user)):
    if data.member_id and not user.get("member_id"):
        member = await db.members.find_one({"id": data.member_id}, {"_id":0})
    else:
        member = await db.members.find_one(own_member_filter(user), {"_id":0})
    if not member: raise HTTPException(404, "Data keanggotaan belum terdaftar. Hubungi petugas perpustakaan.")
    book = await db.books.find_one({"id": data.book_id}, {"_id":0})
    if not book: raise HTTPException(404, "Buku tidak ditemukan")
    if int(book.get("stock") or 0) <= 0: raise HTTPException(400, "Stok buku ini kosong, tidak bisa dipesan")
    if await db.reservations.find_one({"book_id": book["id"], "member_id": member["id"], "status": "Menunggu"}):
        raise HTTPException(400, "Kamu sudah memesan buku ini")
    if await db.reservations.count_documents({"member_id": member["id"], "status": "Menunggu"}) >= 3:
        raise HTTPException(400, "Maksimal 3 pesanan aktif")
    doc = {"id":"RS-"+str(uuid.uuid4())[:8].upper(),"book_id":book["id"],"book_title":book["title"],"book_shelf":book["shelf"],"member_id":member["id"],"member_name":member["name"],"class_name":member["class_name"],"status":"Menunggu","created_at":now_iso(),"processed_at":None,"processed_by":None}
    await db.reservations.insert_one(doc); return public(doc)

@api.delete("/reservations/{res_id}")
async def cancel_reservation(res_id: str, member_id: str = "", user=Depends(current_user)):
    res = await db.reservations.find_one({"id": res_id}, {"_id":0})
    if not res: raise HTTPException(404, "Pesanan tidak ditemukan")
    if user["role"] == "siswa":
        if member_id and not user.get("member_id"):
            if member_id != res["member_id"]: raise HTTPException(403, "Bukan pesanan anggota ini")
        else:
            member = await db.members.find_one(own_member_filter(user), {"_id":0})
            if not member or member["id"] != res["member_id"]: raise HTTPException(403, "Bukan pesanan Anda")
    if res["status"] != "Menunggu": raise HTTPException(400, "Pesanan sudah diproses")
    await db.reservations.update_one({"id": res_id}, {"$set": {"status":"Dibatalkan","processed_at":now_iso(),"processed_by":user["name"]}})
    return {"ok": True}

@api.post("/reservations/{res_id}/process")
async def process_reservation(res_id: str, data: ProcessInput, user=Depends(current_user)):
    staff_only(user)
    res = await db.reservations.find_one({"id": res_id}, {"_id":0})
    if not res or res["status"] != "Menunggu": raise HTTPException(404, "Pesanan menunggu tidak ditemukan")
    if data.action == "pinjam":
        book = await db.books.find_one({"id": res["book_id"]}, {"_id":0})
        if not book: raise HTTPException(404, "Buku tidak ditemukan")
        if await open_loan_count(res["member_id"], book["id"]) > 0: raise HTTPException(400, "Anggota ini masih meminjam buku yang sama")
        await take_copy(book["id"])
        rm = await db.members.find_one({"id": res["member_id"]}, {"_id":0, "class_name":1}) or {}
        await db.transactions.insert_one({"id":"TRX-"+str(uuid.uuid4())[:8].upper(),"member_id":res["member_id"],"member_name":res["member_name"],"class_name":rm.get("class_name",""),"book_id":book["id"],"book_title":book["title"],"action":"pinjam","date":now_iso(),"reservation_id":res_id})
        status = "Dipinjamkan"
    elif data.action == "tolak": status = "Ditolak"
    else: raise HTTPException(400, "Aksi tidak valid")
    await db.reservations.update_one({"id": res_id}, {"$set": {"status":status,"processed_at":now_iso(),"processed_by":user["name"]}})
    return {"ok": True, "status": status}

class StocktakeInput(BaseModel):
    shelf: str
    scanned: List[str] = []
    unscanned: List[str] = []
    misplaced: List[str] = []

@api.get("/stocktake")
async def stocktake_list(user=Depends(current_user)):
    staff_only(user)
    return await db.stocktake.find({}, {"_id":0}).sort("date", -1).to_list(50)

@api.post("/stocktake")
async def stocktake_save(data: StocktakeInput, user=Depends(current_user)):
    staff_only(user)
    doc = {"id":"SO-"+str(uuid.uuid4())[:8].upper(), **data.model_dump(), "scanned_count":len(data.scanned), "unscanned_count":len(data.unscanned), "misplaced_count":len(data.misplaced), "date":now_iso(), "by":user["name"]}
    await db.stocktake.insert_one(doc); return public(doc)

@api.get("/settings")
async def get_settings(user=Depends(current_user)):
    doc = await db.settings.find_one({"id":"main"}, {"_id":0,"id":0})
    return {**DEFAULT_SETTINGS, "member_photo_upload": True, **(doc or {})}

class PhotoToggleInput(BaseModel):
    enabled: bool

@api.put("/settings/photo-upload")
async def put_photo_upload(data: PhotoToggleInput, user=Depends(current_user)):
    staff_only(user)
    await db.settings.update_one({"id":"main"}, {"$set": {"member_photo_upload": data.enabled}}, upsert=True)
    return {"member_photo_upload": data.enabled}

@api.put("/settings")
async def put_settings(data: SettingsInput, user=Depends(current_user)):
    staff_only(user)
    await db.settings.update_one({"id":"main"}, {"$set": data.model_dump()}, upsert=True)
    return data.model_dump()

# ===== Izin unggah foto oleh siswa (diatur admin/petugas) =====
class StudentPhotoConfig(BaseModel):
    enabled: bool = True
    until: str = ""            # ISO datetime; kosong = tanpa batas waktu
    allow_replace: bool = True # boleh mengganti foto yang sudah ada

async def student_photo_config():
    doc = await db.settings.find_one({"id": "student_photo"}, {"_id": 0, "id": 0}) or {}
    cfg = {**StudentPhotoConfig().model_dump(), **doc}
    open_now = bool(cfg["enabled"])
    if open_now and cfg.get("until"):
        try:
            until = datetime.fromisoformat(cfg["until"].replace("Z", "+00:00"))
            if until.tzinfo is None: until = until.replace(tzinfo=timezone.utc)
            if datetime.now(timezone.utc) >= until:
                open_now = False
        except ValueError:
            pass
    cfg["open"] = open_now
    return cfg

async def check_student_photo_allowed(member):
    cfg = await student_photo_config()
    if not cfg["open"]:
        raise HTTPException(403, "Unggah foto sedang ditutup oleh petugas perpustakaan")
    if member and member.get("photo") and not cfg["allow_replace"]:
        raise HTTPException(403, "Foto anggota ini sudah ada. Hubungi petugas untuk menggantinya")

@api.get("/settings/student-photo")
async def get_student_photo(user=Depends(current_user)):
    return await student_photo_config()

@api.put("/settings/student-photo")
async def put_student_photo(data: StudentPhotoConfig, user=Depends(current_user)):
    staff_only(user)
    doc = data.model_dump()
    if doc["until"]:
        try:
            u = datetime.fromisoformat(doc["until"].replace("Z", "+00:00"))
            if u.tzinfo is None: u = u.replace(tzinfo=timezone.utc)
            doc["until"] = u.astimezone(timezone.utc).isoformat()
        except ValueError:
            raise HTTPException(400, "Format batas waktu tidak valid")
    doc["updated_by"] = user["name"]; doc["updated_at"] = now_iso()
    await db.settings.update_one({"id": "student_photo"}, {"$set": doc}, upsert=True)
    return await student_photo_config()

@api.get("/books")
async def books(user=Depends(current_user)):
    rows = await db.books.find({}, {"_id":0}).sort("created_at", -1).to_list(500)
    waiting = {r["_id"]: r["n"] for r in await db.reservations.aggregate([{"$match":{"status":"Menunggu"}},{"$group":{"_id":"$book_id","n":{"$sum":1}}}]).to_list(500)}
    for b in rows: b["waiting"] = waiting.get(b["id"], 0)
    return rows

@api.post("/books")
async def add_book(data: BookInput, user=Depends(current_user)):
    if user["role"] == "siswa": raise HTTPException(403,"Siswa tidak dapat menambah buku")
    same = await find_same_book(data.title, data.author)
    if same:
        await db.books.update_one({"id": same["id"]}, {"$inc": {"stock": data.stock}, **({"$set": {"cover": data.cover}} if data.cover and not same.get("cover") else {})})
        return {**(await refresh_book(same["id"])), "merged": True, "added": data.stock}
    doc = data.model_dump(); doc.update({"id":"BK-"+str(uuid.uuid4())[:8].upper(),"created_at":now_iso(),"borrowed":0,"lost":0}); doc.update(compute_state(doc))
    await db.books.insert_one(doc); return {**public(doc), "merged": False, "added": data.stock}

@api.put("/books/{book_id}")
async def update_book(book_id: str, data: BookUpdate, user=Depends(current_user)):
    staff_only(user)
    book = await db.books.find_one({"id": book_id}, {"_id": 0})
    if not book: raise HTTPException(404, "Buku tidak ditemukan")
    if not await db.shelves.find_one({"id": data.shelf}): raise HTTPException(400, "Rak tidak ditemukan")
    title, author = data.title.strip(), data.author.strip()
    same = await find_same_book(title, author)
    if same and same["id"] != book_id: raise HTTPException(400, f"Judul & penulis ini sudah ada ({same['id']}). Gunakan Atur stok untuk menambah eksemplar.")
    upd = {"title": title, "author": author, "category": data.category.strip() or "Umum", "shelf": data.shelf, "isbn": data.isbn.strip(), "cover": data.cover, "updated_at": now_iso()}
    await db.books.update_one({"id": book_id}, {"$set": upd})
    if title != book["title"]:
        await db.transactions.update_many({"book_id": book_id}, {"$set": {"book_title": title}})
        await db.reservations.update_many({"book_id": book_id}, {"$set": {"book_title": title}})
        await db.missing.update_many({"book_id": book_id}, {"$set": {"title": title}})
    if data.shelf != book.get("shelf"):
        await db.reservations.update_many({"book_id": book_id, "status": "Menunggu"}, {"$set": {"book_shelf": data.shelf}})
    return await refresh_book(book_id)

@api.patch("/books/{book_id}/stock")
async def set_book_stock(book_id: str, data: StockInput, user=Depends(current_user)):
    staff_only(user)
    book = await db.books.find_one({"id": book_id}, {"_id": 0})
    if not book: raise HTTPException(404, "Buku tidak ditemukan")
    if data.stock < int(book.get("borrowed") or 0): raise HTTPException(400, f"Stok tidak boleh kurang dari jumlah yang sedang dipinjam ({book.get('borrowed', 0)})")
    await db.books.update_one({"id": book_id}, {"$set": {"stock": data.stock}})
    return await refresh_book(book_id)

@api.delete("/books/{book_id}")
async def delete_book(book_id: str, user=Depends(current_user)):
    staff_only(user)
    result = await db.books.delete_one({"id":book_id})
    if result.deleted_count == 0: raise HTTPException(404, "Buku tidak ditemukan")
    return {"ok":True}

@api.patch("/books/{book_id}/status")
async def book_status(book_id: str, data: StatusInput, user=Depends(current_user)):
    staff_only(user)
    if data.status not in ["Hilang", "Tersedia"]: raise HTTPException(400, "Status tidak valid")
    book = await db.books.find_one({"id":book_id}, {"_id":0})
    if not book: raise HTTPException(404, "Buku tidak ditemukan")
    if data.status == "Hilang":
        res = await db.books.find_one_and_update({"id": book_id, "$expr": {"$gt": ["$stock", "$borrowed"]}}, {"$inc": {"stock": -1, "lost": 1}})
        if not res: raise HTTPException(400, "Tidak ada eksemplar di rak yang bisa ditandai hilang")
        await refresh_book(book_id)
        report = {"id":"LS-"+str(uuid.uuid4())[:8].upper(),"book_id":book_id,"title":book["title"],"code":book_id,"detail":data.note or "1 eksemplar","date":now_iso(),"status":"Hilang","reported_by":data.reporter.strip() or user["name"]}
        await db.missing.insert_one(report); return public(report)
    res = await db.books.find_one_and_update({"id": book_id, "lost": {"$gt": 0}}, {"$inc": {"stock": 1, "lost": -1}})
    if not res: raise HTTPException(400, "Tidak ada eksemplar buku ini yang tercatat hilang")
    await refresh_book(book_id)
    first = await db.missing.find_one({"book_id": book_id, "status": "Hilang"}, {"_id": 0}, sort=[("date", 1)])
    if first: await db.missing.update_one({"id": first["id"]}, {"$set": {"status": "Ditemukan", "found_at": now_iso()}})
    return {"ok":True, "status":"Tersedia"}

@api.get("/incoming")
async def incoming_list(user=Depends(current_user)):
    staff_only(user)
    return await db.incoming.find({}, {"_id":0}).sort("date",-1).to_list(200)

@api.post("/incoming")
async def add_incoming(data: IncomingInput, user=Depends(current_user)):
    staff_only(user)
    date = data.date or now_iso()[:10]
    same = await find_same_book(data.title, data.author)
    if same:
        await db.books.update_one({"id": same["id"]}, {"$inc": {"stock": data.quantity}, **({"$set": {"cover": data.cover}} if data.cover and not same.get("cover") else {})})
        book = await refresh_book(same["id"]); merged = True
    else:
        base = data.model_dump(exclude={"quantity","date","source"})
        book = {**base, "id":"BK-"+str(uuid.uuid4())[:8].upper(), "stock": data.quantity, "borrowed": 0, "lost": 0, "created_at":now_iso()}; book.update(compute_state(book))
        await db.books.insert_one(book); merged = False
    doc = {"id":"IN-"+str(uuid.uuid4())[:8].upper(),"title":data.title,"author":data.author,"category":data.category,"shelf":data.shelf,"quantity":data.quantity,"detail":f"{data.quantity} eksemplar","source":data.source,"date":date,"status":"Tersimpan","book_ids":[book["id"]],"book_id":book["id"],"merged":merged,"stock_after":book["stock"],"recorded_by":user["name"],"created_at":now_iso()}
    await db.incoming.insert_one(doc); return public(doc)

@api.get("/members")
async def members(user=Depends(current_user)):
    staff_only(user)
    return await db.members.find({}, {"_id":0}).to_list(None)

@api.get("/members/options")
async def member_options(user=Depends(current_user)):
    rows = await db.members.find({}, {"_id":0,"id":1,"name":1,"class_name":1,"student_id":1,"photo":1}).to_list(None)
    rows.sort(key=lambda m: (m.get("class_name",""), m.get("name","")))
    return rows

@api.post("/members")
async def add_member(data: MemberInput, user=Depends(current_user)):
    if user["role"] == "siswa": raise HTTPException(403,"Siswa tidak dapat menambah anggota")
    doc=data.model_dump(); doc.update({"id":"MB-"+str(uuid.uuid4())[:8].upper()}); await db.members.insert_one(doc); return public(doc)

def _cell_text(c):
    if c is None:
        return ""
    if isinstance(c, float) and c.is_integer():
        c = int(c)
    return str(c).strip()

def _parse_member_rows(filename: str, content: bytes):
    name_l = (filename or "").lower()
    rows = []
    if name_l.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(400, "Berkas Excel tidak dapat dibaca. Pastikan format .xlsx yang valid.")
        ws = wb.active
        headers = None
        for row in ws.iter_rows(values_only=True):
            if headers is None:
                cand = [str(c).strip().lower() if c is not None else "" for c in row]
                # Lewati baris judul/petunjuk sampai menemukan baris kolom "nama"
                if not ({"nama", "name", "nama lengkap", "nama_lengkap"} & set(cand)):
                    continue
                headers = cand
                continue
            d = {}
            for i, c in enumerate(row):
                if i < len(headers) and headers[i]:
                    d[headers[i]] = _cell_text(c)
            if any(d.values()):
                rows.append(d)
        if headers is None:
            raise HTTPException(400, "Kolom 'nama' tidak ditemukan. Gunakan template Excel dari tombol Unduh template.")
    elif name_l.endswith(".xls"):
        raise HTTPException(400, "Format .xls lama belum didukung. Di Excel pilih Simpan Sebagai → Excel Workbook (.xlsx).")
    else:
        raise HTTPException(400, "Format tidak didukung. Gunakan berkas Excel (.xlsx)")
    return rows

@api.get("/members/import-template.xlsx")
async def member_import_template(user=Depends(current_user)):
    staff_only(user)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    wb = Workbook(); ws = wb.active; ws.title = "Anggota"
    thin = Side(style="thin", color="CBD5E1"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    headers = ["nama", "nis", "kelas"]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F2745")
        c.alignment = Alignment(horizontal="center", vertical="center"); c.border = border
    samples = [("Budi Santoso", "24100", "X IPA 1"), ("Siti Aminah", "24101", "X IPA 1"), ("Rudi Hartono", "24102", "X IPS 2")]
    for i, r in enumerate(samples, start=2):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v); c.border = border
            if j == 2: c.number_format = "@"
    for row in range(2, 2002):
        ws.cell(row=row, column=2).number_format = "@"  # NIS sebagai teks (angka 0 di depan tidak hilang)
    ws.column_dimensions["A"].width = 32; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 14
    ws.row_dimensions[1].height = 22; ws.freeze_panes = "A2"
    info = wb.create_sheet("Petunjuk")
    for i, t in enumerate([
        "PETUNJUK IMPOR ANGGOTA — PERPUSTAKAAN SMA NEGERI 1 ABUNG TENGAH",
        "1. Isi data di lembar 'Anggota'. Jangan ubah judul kolom: nama, nis, kelas.",
        "2. Satu baris = satu anggota. Boleh banyak kelas dalam satu berkas.",
        "3. Hapus 3 baris contoh (Budi, Siti, Rudi) sebelum mengimpor.",
        "4. NIS yang sudah terdaftar akan dilewati otomatis.",
        "5. Simpan sebagai .xlsx lalu unggah di menu Anggota → Impor.",
    ], start=1):
        c = info.cell(row=i, column=1, value=t)
        if i == 1: c.font = Font(bold=True, color="0F2745", size=12)
    info.column_dimensions["A"].width = 80
    buf = io.BytesIO(); wb.save(buf)
    return Response(content=buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="template-anggota.xlsx"'})

def _pick(d, *keys):
    for k in keys:
        if d.get(k):
            return d[k]
    return ""

@api.post("/members/import")
async def import_members(file: UploadFile = File(...), user=Depends(current_user)):
    staff_only(user)
    content = await file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(400, "Ukuran berkas maksimal 5MB")
    rows = _parse_member_rows(file.filename or "", content)
    if not rows:
        raise HTTPException(400, "Berkas kosong atau tidak memiliki data")
    if len(rows) > 2000:
        raise HTTPException(400, "Maksimal 2000 baris per impor")
    existing = {m.get("student_id") for m in await db.members.find({}, {"_id": 0, "student_id": 1}).to_list(None) if m.get("student_id")}
    seen = set()
    to_insert, errors, skipped = [], [], 0
    for idx, d in enumerate(rows, start=2):
        name = _pick(d, "nama", "name", "nama lengkap", "nama_lengkap")
        sid = _pick(d, "nis", "student_id", "nomor induk", "no induk", "nomor_induk", "no_induk")
        cls = _pick(d, "kelas", "class_name", "class", "rombel")
        reason = "Nama atau kelas kosong" if not name or not cls else (f"NIS {sid} sudah terdaftar" if sid and (sid in existing or sid in seen) else None)
        if reason:
            skipped += 1
            errors.append({"row": idx, "nama": name, "nis": sid, "kelas": cls, "alasan": reason})
            continue
        if sid:
            seen.add(sid)
        to_insert.append({"id": "MB-" + str(uuid.uuid4())[:8].upper(), "name": name, "student_id": sid, "class_name": cls, "photo": ""})
    if to_insert:
        await db.members.insert_many(to_insert)
    return {"imported": len(to_insert), "skipped": skipped, "total": len(rows), "errors": errors}

class FailedRow(BaseModel):
    row: int
    nama: str = ""
    nis: str = ""
    kelas: str = ""
    alasan: str = ""

@api.post("/members/import-failed.xlsx")
async def import_failed_xlsx(rows: List[FailedRow], user=Depends(current_user)):
    staff_only(user)
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(); ws = wb.active; ws.title = "Gagal diimpor"
    ws.append(["nama", "nis", "kelas", "alasan", "baris asal"])
    for c in ws[1]: c.font = Font(bold=True)
    for r in rows: ws.append([r.nama, r.nis, r.kelas, r.alasan, r.row])
    for col, w in zip("ABCDE", (32, 16, 12, 36, 11)): ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="anggota-gagal-impor.xlsx"'})

@api.put("/members/{member_id}")
async def update_member(member_id: str, data: MemberInput, user=Depends(current_user)):
    staff_only(user)
    member = await db.members.find_one({"id": member_id}, {"_id":0})
    if not member: raise HTTPException(404, "Anggota tidak ditemukan")
    await db.members.update_one({"id": member_id}, {"$set": data.model_dump()})
    updated = await db.members.find_one({"id": member_id}, {"_id":0})
    return updated

@api.post("/members/{member_id}/photo")
async def set_member_photo(member_id: str, data: PhotoInput, user=Depends(current_user)):
    await check_student_photo(user, data.photo, member_id)
    member = await db.members.find_one({"id": member_id}, {"_id":0})
    if not member: raise HTTPException(404, "Anggota tidak ditemukan")
    if user["role"] == "siswa":
        await check_student_photo_allowed(member)
    await db.members.update_one({"id": member_id}, {"$set": {"photo": data.photo}})
    return await db.members.find_one({"id": member_id}, {"_id":0})

@api.post("/members/promote")
async def promote_members(data: PromoteInput, user=Depends(current_user)):
    staff_only(user)
    from_c = data.from_class.strip()
    to_c = data.to_class.strip()
    if not from_c or not to_c:
        raise HTTPException(400, "Kelas asal dan kelas tujuan wajib diisi")
    if from_c == to_c:
        raise HTTPException(400, "Kelas asal dan tujuan tidak boleh sama")
    count = await db.members.count_documents({"class_name": from_c})
    if count == 0:
        raise HTTPException(404, "Tidak ada anggota pada kelas asal tersebut")
    result = await db.members.update_many({"class_name": from_c}, {"$set": {"class_name": to_c}})
    return {"ok": True, "moved": result.modified_count, "from_class": from_c, "to_class": to_c}

@api.post("/members/bulk-delete")
async def bulk_delete_members(data: BulkDeleteInput, user=Depends(current_user)):
    staff_only(user)
    members = await db.members.find({"class_name": data.class_name}, {"_id":0,"id":1}).to_list(None)
    ids = [m["id"] for m in members]
    if not ids: raise HTTPException(404, "Tidak ada anggota pada kelas tersebut")
    await db.reservations.update_many({"member_id": {"$in": ids}, "status": "Menunggu"}, {"$set": {"status":"Dibatalkan","processed_at":now_iso(),"processed_by":user["name"]}})
    result = await db.members.delete_many({"id": {"$in": ids}})
    await db.users.delete_many({"member_id": {"$in": ids}, "role": "siswa"})
    return {"ok": True, "deleted": result.deleted_count, "class_name": data.class_name}

@api.delete("/members/{member_id}")
async def delete_member(member_id: str, user=Depends(current_user)):
    staff_only(user)
    result = await db.members.delete_one({"id": member_id})
    if result.deleted_count == 0: raise HTTPException(404, "Anggota tidak ditemukan")
    await db.users.delete_many({"member_id": member_id, "role": "siswa"})
    await db.reservations.update_many({"member_id": member_id, "status": "Menunggu"}, {"$set": {"status":"Dibatalkan","processed_at":now_iso(),"processed_by":user["name"]}})
    return {"ok": True}

def _mask_email(e):
    name, _, dom = e.partition("@")
    return f"{name[:3]}{'*' * max(len(name) - 3, 1)}@{dom}"

CODE_PURPOSES = {"reset": "reset data", "restore": "pemulihan data dari cadangan"}

@api.post("/admin/reset-code")
async def send_reset_code(purpose: str = "reset", user=Depends(current_user)):
    if user["role"] != "admin": raise HTTPException(403, "Hanya admin yang dapat mereset atau memulihkan data")
    if purpose not in CODE_PURPOSES: raise HTTPException(400, "Jenis kode tidak dikenal")
    action = CODE_PURPOSES[purpose]
    if not os.environ.get("EMERGENT_EMAIL_KEY"): raise HTTPException(503, "Layanan email belum dikonfigurasi (EMERGENT_EMAIL_KEY belum diatur)")
    from mailer import send_email
    from html import escape
    await db.reset_codes.update_one({"id": purpose}, {"$setOnInsert": {"id": purpose}}, upsert=True)
    cutoff = (datetime.now(timezone.utc) - timedelta(seconds=60)).isoformat()
    last = await db.reset_codes.find_one_and_update({"id": purpose, "sent_at": {"$not": {"$gt": cutoff}}}, {"$set": {"sent_at": datetime.now(timezone.utc).isoformat()}})
    if last is None:
        raise HTTPException(429, "Tunggu 1 menit sebelum meminta kode baru")
    code = f"{secrets.randbelow(1000000):06d}"
    to = os.environ["RESET_VERIFY_EMAIL"]
    html = ('<table role="presentation" width="100%"><tr><td style="padding:24px;font-family:Arial,sans-serif;color:#1f1640">'
            f'<p style="font-size:16px;margin:0 0 12px">Permintaan <b>{action}</b> perpustakaan oleh '
            f'{escape(user["name"])}.</p><p style="margin:0 0 8px">Kode verifikasi Anda:</p>'
            f'<p style="font-size:32px;letter-spacing:8px;font-weight:bold;color:#6d28d9;margin:0 0 12px">{code}</p>'
            '<p style="margin:0 0 12px">Kode berlaku 10 menit dan hanya bisa dipakai sekali. Jika Anda tidak memintanya, abaikan email ini dan jangan bagikan kode kepada siapa pun.</p>'
            f'<p style="font-size:12px;color:#888">Dikirim oleh {escape(os.environ["EMAIL_FROM_NAME"])}.</p></td></tr></table>')
    try:
        await send_email(to=to, subject=f"Kode verifikasi {action} perpustakaan", html=html)
    except Exception:
        await db.reset_codes.update_one({"id": purpose}, {"$set": {"sent_at": last.get("sent_at", "")}})
        raise
    await db.reset_codes.update_one({"id": purpose}, {"$set": {"id": purpose, "code_hash": bcrypt.hashpw(code.encode(), bcrypt.gensalt()).decode(),
        "expires_at": (datetime.now(timezone.utc) + timedelta(minutes=10)).isoformat(), "sent_at": datetime.now(timezone.utc).isoformat(), "attempts": 0}}, upsert=True)
    return {"ok": True, "sent_to": _mask_email(to)}

async def verify_reset_code(code: str, purpose: str = "reset"):
    rec = await db.reset_codes.find_one({"id": purpose}, {"_id": 0})
    if not rec or rec["expires_at"] < datetime.now(timezone.utc).isoformat():
        raise HTTPException(400, "Kode verifikasi tidak ada atau sudah kedaluwarsa. Minta kode baru.")
    if rec.get("attempts", 0) >= 5:
        await db.reset_codes.delete_one({"id": purpose})
        raise HTTPException(400, "Terlalu banyak percobaan. Minta kode baru.")
    if not code.strip() or not bcrypt.checkpw(code.strip().encode(), rec["code_hash"].encode()):
        await db.reset_codes.update_one({"id": purpose}, {"$inc": {"attempts": 1}})
        raise HTTPException(400, "Kode verifikasi salah")
    await db.reset_codes.delete_one({"id": purpose})

@api.post("/admin/reset-data")
async def reset_data(data: ResetInput, user=Depends(current_user)):
    if user["role"] != "admin": raise HTTPException(403, "Hanya admin yang dapat mereset data")
    await verify_reset_code(data.code)
    members = (await db.members.delete_many({})).deleted_count
    await db.users.delete_many({"role": "siswa", "member_id": {"$exists": True}})
    transactions = (await db.transactions.delete_many({})).deleted_count
    incoming = (await db.incoming.delete_many({})).deleted_count
    missing = (await db.missing.delete_many({})).deleted_count
    reservations = (await db.reservations.delete_many({})).deleted_count
    stocktake = (await db.stocktake.delete_many({})).deleted_count
    if data.include_books:
        books = (await db.books.delete_many({})).deleted_count
    else:
        books = 0
        await db.books.update_many({}, [{"$set": {"stock": {"$add": [{"$ifNull": ["$stock", 0]}, {"$ifNull": ["$lost", 0]}]}, "lost": 0, "borrowed": 0}}, {"$set": {"available": "$stock", "status": {"$cond": [{"$gt": ["$stock", 0]}, "Tersedia", "Habis"]}}}])
    await db.visits.delete_many({})
    await db.meta.update_one({"id": "seed"}, {"$set": {"id": "seed", "done": True}}, upsert=True)
    return {"ok": True, "deleted": {"members": members, "transactions": transactions, "incoming": incoming, "missing": missing, "reservations": reservations, "stocktake": stocktake, "books": books}, "books_deleted": data.include_books}

BACKUP_COLLECTIONS = ["books", "members", "transactions", "incoming", "missing", "reservations", "stocktake", "shelves", "settings", "visits"]

@api.get("/admin/backup")
async def backup_data(user=Depends(current_user)):
    if user["role"] != "admin": raise HTTPException(403, "Hanya admin yang dapat mengunduh cadangan data")
    dump = {}
    for name in BACKUP_COLLECTIONS:
        dump[name] = await db[name].find({}, {"_id": 0}).to_list(100000)
    return {"app": "perpustakaan-sma1-abung-tengah", "version": 1, "generated_at": now_iso(), "collections": dump}

BACKUP_REPORTS = [("transaksi", "01-Laporan-Transaksi"), ("masuk", "02-Laporan-Buku-Masuk"), ("hilang", "03-Laporan-Buku-Hilang"), ("kunjungan", "04-Laporan-Pengunjung")]

@api.get("/admin/backup.zip")
async def backup_zip(user=Depends(current_user)):
    """Cadangan lengkap: data (JSON, bisa dipulihkan) + semua laporan siap baca (Excel & PDF)."""
    if user["role"] != "admin": raise HTTPException(403, "Hanya admin yang dapat mengunduh cadangan data")
    import zipfile, asyncio
    payload = await backup_data(user)
    now = datetime.now(WIB); stamp = now.strftime("%Y-%m-%d")
    jobs, notes = [], []
    for kind, name in BACKUP_REPORTS:
        rows, settings, _ = await export_rows(kind, "semua")
        jobs.append((f"laporan/{name}", kind, "semua", rows, settings, len(rows)))
    lb_months = (await visit_leaderboard(None, limit=1))["months"]
    for month in lb_months:
        lb, settings, _ = await export_rows("peringkat", month)
        if lb["total_visits"] == 0 and month != now.strftime("%Y-%m"):
            continue
        jobs.append((f"laporan/05-Peringkat-Rajin-Berkunjung/Peringkat-{month}", "peringkat", lb["label"], lb, settings, len(lb["students"])))
    for month in (await loan_leaderboard(None, limit=1))["months"]:
        lb, settings, _ = await export_rows("peminjam", month)
        if lb["total_loans"] == 0 and month != now.strftime("%Y-%m"):
            continue
        jobs.append((f"laporan/06-Peringkat-Rajin-Meminjam/Peringkat-{month}", "peminjam", lb["label"], lb, settings, len(lb["students"])))

    def build():
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("data-cadangan.json", json.dumps(payload, ensure_ascii=False, indent=2, default=str))
            for path, kind, period, rows, settings, count in jobs:
                for ext, render in (("xlsx", render_report_xlsx), ("pdf", render_report_pdf)):
                    try:
                        z.writestr(f"{path}.{ext}", render(kind, period, rows, settings, user["name"]))
                    except Exception as e:
                        logging.getLogger(__name__).warning("Laporan %s.%s gagal dibuat: %s", path, ext, e)
                        notes.append(f"- {path}.{ext} gagal dibuat")
                notes_line = f"- {path}.xlsx / .pdf ({count} baris)"
                notes.append(notes_line)
            readme = [
                "CADANGAN PERPUSTAKAAN SMA NEGERI 1 ABUNG TENGAH",
                f"Dibuat: {fmt_long(now, True)} WIB oleh {user['name']}",
                "",
                "data-cadangan.json  -> berkas untuk menu 'Pulihkan data' (boleh juga unggah ZIP ini langsung).",
                "laporan/            -> laporan siap baca (Excel & PDF), semua periode:",
                *notes,
            ]
            z.writestr("BACA-SAYA.txt", "\n".join(readme))
        return buf.getvalue()

    data = await asyncio.to_thread(build)
    return Response(content=data, media_type="application/zip", headers={"Content-Disposition": f'attachment; filename="backup-perpustakaan-{stamp}.zip"'})

@api.post("/admin/restore")
async def restore_data(file: UploadFile = File(...), code: str = Form(""), user=Depends(current_user)):
    if user["role"] != "admin": raise HTTPException(403, "Hanya admin yang dapat memulihkan data")
    await verify_reset_code(code, "restore")
    content = await file.read()
    if len(content) > 50 * 1024 * 1024:
        raise HTTPException(400, "Ukuran berkas cadangan maksimal 50MB")
    if content[:2] == b"PK":  # berkas ZIP dari "Unduh cadangan"
        import zipfile
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                content = z.read("data-cadangan.json")
        except KeyError:
            raise HTTPException(400, "ZIP cadangan tidak berisi data-cadangan.json")
        except zipfile.BadZipFile:
            raise HTTPException(400, "Berkas ZIP cadangan rusak")
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(400, "Ukuran data cadangan maksimal 20MB")
    try:
        payload = json.loads(content.decode("utf-8-sig", errors="replace"))
    except Exception:
        raise HTTPException(400, "Berkas cadangan tidak valid (bukan JSON)")
    collections = payload.get("collections")
    if not isinstance(collections, dict):
        raise HTTPException(400, "Struktur berkas cadangan tidak dikenali")
    restored = {}
    for name in BACKUP_COLLECTIONS:
        rows = collections.get(name)
        if rows is None:
            continue
        if not isinstance(rows, list):
            raise HTTPException(400, f"Data '{name}' pada cadangan tidak valid")
        await db[name].delete_many({})
        clean = [{k: v for k, v in r.items() if k != "_id"} for r in rows if isinstance(r, dict)]
        if clean:
            await db[name].insert_many(clean)
        restored[name] = len(clean)
    # Cadangan format lama (buku per eksemplar tanpa angka stok) digabung menjadi satu judul
    if await db.books.count_documents({"stock": {"$exists": False}}) > 0:
        await normalize_books()
    await db.meta.update_one({"id": "seed"}, {"$set": {"id": "seed", "done": True}}, upsert=True)
    return {"ok": True, "restored": restored}


@api.get("/members/{member_id}/history")
async def member_history(member_id: str, user=Depends(current_user)):
    staff_only(user)
    member = await db.members.find_one({"id":member_id}, {"_id":0})
    if not member: raise HTTPException(404, "Anggota tidak ditemukan")
    return await history_for(member)

@api.get("/account/history")
async def my_history(user=Depends(current_user)):
    if user["role"] != "siswa" or not user.get("member_id"): raise HTTPException(403, "Riwayat hanya untuk akun anggota")
    member = await db.members.find_one({"id": user["member_id"]}, {"_id":0})
    if not member: raise HTTPException(404, "Data anggota tidak ditemukan")
    return await history_for(member)

class MemberAccountInput(BaseModel):
    email: str = Field(..., min_length=3, max_length=120)
    password: str = Field("", max_length=100)

@api.get("/members/{member_id}/account")
async def get_member_account(member_id: str, user=Depends(current_user)):
    staff_only(user)
    acc = await db.users.find_one({"member_id": member_id, "role": "siswa"}, {"_id":0, "password_hash":0})
    return acc

@api.put("/members/{member_id}/account")
async def set_member_account(member_id: str, data: MemberAccountInput, user=Depends(current_user)):
    staff_only(user)
    member = await db.members.find_one({"id": member_id}, {"_id":0})
    if not member: raise HTTPException(404, "Anggota tidak ditemukan")
    email = data.email.strip().lower()
    if not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email): raise HTTPException(400, "Format email tidak valid")
    acc = await db.users.find_one({"member_id": member_id, "role": "siswa"}, {"_id":0})
    other = await db.users.find_one({"email": email}, {"_id":0})
    if other and (not acc or other["id"] != acc["id"]): raise HTTPException(400, "Email sudah dipakai akun lain")
    if data.password and len(data.password) < 6: raise HTTPException(400, "Kata sandi minimal 6 karakter")
    upd = {"email": email, "name": member["name"], "student_id": member.get("student_id", ""), "class_name": member.get("class_name", "")}
    if data.password: upd["password_hash"] = password_hash(data.password)
    if acc:
        await db.users.update_one({"id": acc["id"]}, {"$set": upd})
    else:
        if not data.password: raise HTTPException(400, "Kata sandi wajib diisi untuk akun baru")
        await db.users.insert_one({"id": "u-" + str(uuid.uuid4())[:8], "role": "siswa", "member_id": member_id, **upd})
    return await db.users.find_one({"member_id": member_id, "role": "siswa"}, {"_id":0, "password_hash":0})

@api.delete("/members/{member_id}/account")
async def delete_member_account(member_id: str, user=Depends(current_user)):
    staff_only(user)
    res = await db.users.delete_one({"member_id": member_id, "role": "siswa"})
    if res.deleted_count == 0: raise HTTPException(404, "Anggota ini belum punya akun")
    return {"ok": True}

async def history_for(member):
    member_id = member["id"]
    tx = await db.transactions.find({"member_id":member_id}, {"_id":0}).sort("date",1).to_list(300)
    history = []
    for t in tx:
        if t["action"] == "pinjam": history.append({**t, "returned": False, "returned_at": None})
        else:
            for h in reversed(history):
                if h["book_id"] == t["book_id"] and not h["returned"]: h.update(returned=True, returned_at=t["date"]); break
    history.reverse()
    active = sum(1 for h in history if not h["returned"])
    return {"member":member,"history":history,"total":len(history),"active":active}

@api.get("/shelves")
async def shelves(user=Depends(current_user)): return await db.shelves.find({}, {"_id":0}).sort("id", 1).to_list(200)

@api.post("/shelves")
async def add_shelf(data: ShelfInput, user=Depends(current_user)):
    staff_only(user)
    sid = data.id.strip().upper()
    if await db.shelves.find_one({"id": sid}): raise HTTPException(400, f"Kode rak {sid} sudah dipakai")
    doc = {"id": sid, "name": data.name.strip(), "capacity": data.capacity, "color": data.color, "created_at": now_iso()}
    await db.shelves.insert_one(doc); return public(doc)

@api.put("/shelves/{shelf_id}")
async def update_shelf(shelf_id: str, data: ShelfInput, user=Depends(current_user)):
    staff_only(user)
    shelf = await db.shelves.find_one({"id": shelf_id}, {"_id": 0})
    if not shelf: raise HTTPException(404, "Rak tidak ditemukan")
    new_id = data.id.strip().upper()
    if new_id != shelf_id and await db.shelves.find_one({"id": new_id}): raise HTTPException(400, f"Kode rak {new_id} sudah dipakai")
    await db.shelves.update_one({"id": shelf_id}, {"$set": {"id": new_id, "name": data.name.strip(), "capacity": data.capacity, "color": data.color, "updated_at": now_iso()}})
    moved = 0
    if new_id != shelf_id:
        moved = (await db.books.update_many({"shelf": shelf_id}, {"$set": {"shelf": new_id}})).modified_count
        await db.incoming.update_many({"shelf": shelf_id}, {"$set": {"shelf": new_id}})
        await db.stocktake.update_many({"shelf": shelf_id}, {"$set": {"shelf": new_id}})
        await db.reservations.update_many({"book_shelf": shelf_id}, {"$set": {"book_shelf": new_id}})
    return {**(await db.shelves.find_one({"id": new_id}, {"_id": 0})), "books_updated": moved}

@api.delete("/shelves/{shelf_id}")
async def delete_shelf(shelf_id: str, move_to: str = "", user=Depends(current_user)):
    staff_only(user)
    if not await db.shelves.find_one({"id": shelf_id}): raise HTTPException(404, "Rak tidak ditemukan")
    if await db.shelves.count_documents({}) <= 1: raise HTTPException(400, "Minimal harus ada satu rak")
    count = await db.books.count_documents({"shelf": shelf_id})
    if count:
        if not move_to: raise HTTPException(400, f"Rak masih berisi {count} judul buku. Pilih rak tujuan untuk memindahkannya.")
        if move_to == shelf_id or not await db.shelves.find_one({"id": move_to}): raise HTTPException(400, "Rak tujuan tidak valid")
        await db.books.update_many({"shelf": shelf_id}, {"$set": {"shelf": move_to}})
        await db.reservations.update_many({"book_shelf": shelf_id, "status": "Menunggu"}, {"$set": {"book_shelf": move_to}})
    await db.shelves.delete_one({"id": shelf_id})
    return {"ok": True, "moved": count, "move_to": move_to or None}

@api.get("/transactions")
async def transactions(user=Depends(current_user)):
    staff_only(user)
    return await db.transactions.find({}, {"_id":0}).sort("date",-1).to_list(200)

@api.post("/transactions")
async def transaction(data: TransactionInput, user=Depends(current_user)):
    staff_only(user)
    member=await db.members.find_one({"id":data.member_id},{"_id":0}); book=await db.books.find_one({"id":data.book_id},{"_id":0})
    if not member or not book: raise HTTPException(404,"Anggota atau buku tidak ditemukan")
    action=data.action
    if action not in ("pinjam","kembali"): raise HTTPException(400,"Aksi tidak valid")
    open_count = await open_loan_count(member["id"], book["id"])
    if action=="pinjam":
        if open_count > 0: raise HTTPException(400,"Anggota ini masih meminjam buku yang sama")
        await take_copy(book["id"])
    else:
        if open_count <= 0: raise HTTPException(400,"Anggota ini tidak sedang meminjam buku ini")
        await return_copy(book["id"])
    doc={"id":"TRX-"+str(uuid.uuid4())[:8].upper(),"member_id":member["id"],"member_name":member["name"],"class_name":member.get("class_name",""),"book_id":book["id"],"book_title":book["title"],"action":action,"date":now_iso()}; await db.transactions.insert_one(doc); return public(doc)

@api.get("/reports")
async def reports(user=Depends(current_user)):
    staff_only(user)
    tx=await db.transactions.find({}, {"_id":0}).sort("date",-1).to_list(500)
    incoming=await db.incoming.find({}, {"_id":0}).sort("date",-1).to_list(200)
    missing=await db.missing.find({}, {"_id":0}).sort("date",-1).to_list(200)
    return {"transactions":tx,"loans":await attach_class(build_loans(tx)),"incoming":incoming,"missing":missing}

async def attach_class(loans):
    """Lengkapi kelas peminjam: pakai kelas saat transaksi, atau kelas anggota sekarang untuk transaksi lama."""
    ids = list({l["member_id"] for l in loans if not l.get("class_name") and l.get("member_id")})
    info = {m["id"]: m.get("class_name", "") for m in await db.members.find({"id": {"$in": ids}}, {"_id":0, "id":1, "class_name":1}).to_list(100000)} if ids else {}
    for l in loans:
        if not l.get("class_name"): l["class_name"] = info.get(l.get("member_id"), "")
    return loans

async def loan_leaderboard(month: Optional[str] = None, limit: int = 10):
    """Peringkat siswa & kelas dengan peminjaman terbanyak dalam satu bulan (WIB)."""
    now = datetime.now(WIB)
    if month and not re.match(r"^\d{4}-(0[1-9]|1[0-2])$", month): raise HTTPException(400, "Format bulan harus YYYY-MM")
    key = month or now.strftime("%Y-%m")
    y, m = (int(x) for x in key.split("-"))
    loans = []
    for t in await db.transactions.find({"action": "pinjam"}, {"_id": 0}).to_list(100000):
        try: t["_wib"] = datetime.fromisoformat(t["date"].replace("Z", "+00:00")).astimezone(WIB).strftime("%Y-%m")
        except (ValueError, AttributeError): continue
        loans.append(t)
    rows = await attach_class([t for t in loans if t["_wib"] == key])
    members = {x["id"]: x for x in await db.members.find({"id": {"$in": list({r["member_id"] for r in rows})}}, {"_id": 0, "id": 1, "name": 1, "class_name": 1, "photo": 1}).to_list(None)}
    students, classes = {}, {}
    for r in rows:
        mem = members.get(r["member_id"]) or {}
        st = students.setdefault(r["member_id"], {"member_id": r["member_id"], "name": mem.get("name") or r.get("member_name", "-"), "class_name": mem.get("class_name") or r.get("class_name") or "-", "photo": mem.get("photo", ""), "loans": 0, "titles": set()})
        st["loans"] += 1; st["titles"].add(r.get("book_id"))
        cls = st["class_name"] if st["class_name"] != "-" else "Tanpa kelas"
        c = classes.setdefault(cls, {"class_name": cls, "loans": 0, "members": set()})
        c["loans"] += 1; c["members"].add(r["member_id"])
    total = len(rows)
    pct = lambda n: round(n * 100 / total, 1) if total else 0.0
    ranked = sorted(students.values(), key=lambda s: (-s["loans"], -len(s["titles"]), s["name"]))[:limit]
    out_students = [{**{k: v for k, v in s.items() if k != "titles"}, "rank": i, "titles": len(s["titles"]), "pct": pct(s["loans"])} for i, s in enumerate(ranked, 1)]
    out_classes = [{"rank": i, "class_name": c["class_name"], "loans": c["loans"], "members": len(c["members"]), "pct": pct(c["loans"])} for i, c in enumerate(sorted(classes.values(), key=lambda c: (-c["loans"], -len(c["members"]), c["class_name"])), 1)]
    months = sorted({t["_wib"] for t in loans} | {now.strftime("%Y-%m")}, reverse=True)
    return {"month": key, "label": f"{MONTHS_ID[m - 1]} {y}", "total_loans": total, "active_students": len(students), "students": out_students, "classes": out_classes, "months": months}

@api.get("/reports/loan-leaderboard")
async def get_loan_leaderboard(month: Optional[str] = None, user=Depends(current_user)):
    staff_only(user)
    return await loan_leaderboard(month)

async def export_rows(type: str, period: str):
    if type not in ("transaksi","masuk","hilang","kunjungan","statistik","peringkat","peminjam"): raise HTTPException(400, "Jenis laporan tidak valid")
    now = datetime.now(WIB)
    settings = await db.settings.find_one({"id":"main"}, {"_id":0}) or {}
    if type == "peminjam":
        return await loan_leaderboard(period if re.match(r"^\d{4}-\d{2}$", period or "") else None, limit=50), settings, now
    if type == "peringkat":
        await visits_auto_close()
        return await visit_leaderboard(period if re.match(r"^\d{4}-\d{2}$", period or "") else None, limit=50), settings, now
    if type in ("kunjungan","statistik"):
        await visits_auto_close()
        if type == "statistik":
            return await visit_stats(period), settings, now
        rows = await db.visits.find({}, {"_id":0}).sort("check_in",-1).to_list(100000)
        return [r for r in rows if in_visit_period(r.get("check_in"), period, now)], settings, now
    if type == "transaksi":
        rows = await attach_class(build_loans(await db.transactions.find({}, {"_id":0}).to_list(100000)))
    elif type == "masuk":
        rows = await db.incoming.find({}, {"_id":0}).sort("date",-1).to_list(100000)
    else:
        rows = await db.missing.find({}, {"_id":0}).sort("date",-1).to_list(100000)
    rows = [r for r in rows if in_period(r.get("date"), period, now)]
    return rows, settings, now

@api.get("/reports/export.pdf")
async def export_report(type: str = "transaksi", period: str = "semua", user=Depends(current_user)):
    staff_only(user)
    rows, settings, now = await export_rows(type, period)
    pdf = render_report_pdf(type, rows["label"] if type in ("peringkat", "peminjam") else period, rows, settings, user["name"])
    filename = f"laporan-{type}-{period}-{now.strftime('%Y%m%d')}.pdf"
    return Response(content=pdf, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{filename}"'})

@api.get("/reports/export.xlsx")
async def export_report_xlsx(type: str = "transaksi", period: str = "semua", user=Depends(current_user)):
    staff_only(user)
    rows, settings, now = await export_rows(type, period)
    data = render_report_xlsx(type, rows["label"] if type in ("peringkat", "peminjam") else period, rows, settings, user["name"])
    filename = f"laporan-{type}-{period}-{now.strftime('%Y%m%d')}.xlsx"
    return Response(content=data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})

from visits import build_router, visits_auto_close, visit_stats, in_visit_period, visit_leaderboard, MONTHS_ID
api.include_router(build_router(db, current_user, staff_only, now_iso))
from summary import build_router as summary_router
api.include_router(media_router(db, current_user))
api.include_router(summary_router(db, current_user, staff_only, overdue_loans, visits_auto_close))
app.include_router(api)
_cors = [o.strip() for o in (os.environ.get("FRONTEND_URL", "") + "," + os.environ.get("CORS_ORIGINS", "")).split(",") if o.strip() and o.strip() != "*"]
# Frontend di perpusmansa.vercel.app (proxy /api) atau preview Emergent; CORS hanya untuk asal tambahan
app.add_middleware(CORSMiddleware, allow_credentials=True, allow_origins=_cors, allow_origin_regex=r"https://([a-z0-9-]+\.)*(emergentagent\.com|emergent\.host|vercel\.app)", allow_methods=["*"], allow_headers=["*"])

@app.on_event("shutdown")
async def shutdown(): client.close()