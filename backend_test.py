#!/usr/bin/env python3
"""
Comprehensive Backend API Testing for School Library Migration
Tests all features after migration from backup ZIP
"""

import requests
import sys
import io
import zipfile
import json
from datetime import datetime

BASE_URL = "https://restore-app-34--3g9ynt9t.preview.emergentagent.com/api"

class TestRunner:
    def __init__(self):
        self.tests_run = 0
        self.tests_passed = 0
        self.admin_token = None
        self.petugas_token = None
        self.siswa_token = None
        self.test_book_id = None
        self.test_member_id = None
        self.test_shelf_id = None
        self.test_visit_id = None

    def log(self, msg):
        print(f"  {msg}")

    def test(self, name, fn):
        """Run a single test"""
        self.tests_run += 1
        print(f"\n🔍 Test {self.tests_run}: {name}")
        try:
            fn()
            self.tests_passed += 1
            print(f"✅ PASSED")
            return True
        except AssertionError as e:
            print(f"❌ FAILED: {e}")
            return False
        except Exception as e:
            print(f"❌ ERROR: {e}")
            return False

    # ===== AUTH TESTS =====
    def test_admin_login(self):
        """Test admin login"""
        def _test():
            resp = requests.post(f"{BASE_URL}/auth/login", json={
                "email": "admin@smansa.sch.id",
                "password": "admin123",
                "portal": "admin"
            })
            assert resp.status_code == 200, f"Admin login failed: {resp.status_code} {resp.text}"
            data = resp.json()
            assert data["role"] == "admin", f"Expected admin role, got {data.get('role')}"
            self.admin_token = resp.cookies.get("access_token")
            self.log(f"Admin logged in: {data['name']}")
        self.test("Admin login successful", _test)

    def test_petugas_login(self):
        """Test petugas login"""
        def _test():
            resp = requests.post(f"{BASE_URL}/auth/login", json={
                "email": "petugas@smansa.sch.id",
                "password": "petugas123",
                "portal": "admin"
            })
            assert resp.status_code == 200, f"Petugas login failed: {resp.status_code} {resp.text}"
            data = resp.json()
            assert data["role"] == "petugas", f"Expected petugas role, got {data.get('role')}"
            self.petugas_token = resp.cookies.get("access_token")
            self.log(f"Petugas logged in: {data['name']}")
        self.test("Petugas login successful", _test)

    def test_siswa_login_siswa_portal(self):
        """Test siswa login at siswa portal"""
        def _test():
            resp = requests.post(f"{BASE_URL}/auth/login", json={
                "email": "siswa@smansa.sch.id",
                "password": "siswa123",
                "portal": "siswa"
            })
            assert resp.status_code == 200, f"Siswa login failed: {resp.status_code} {resp.text}"
            data = resp.json()
            assert data["role"] == "siswa", f"Expected siswa role, got {data.get('role')}"
            self.siswa_token = resp.cookies.get("access_token")
            self.log(f"Siswa logged in: {data['name']}")
        self.test("Siswa login at /siswa portal successful", _test)

    def test_siswa_rejected_at_admin_portal(self):
        """Test siswa login rejected at admin portal"""
        def _test():
            resp = requests.post(f"{BASE_URL}/auth/login", json={
                "email": "siswa@smansa.sch.id",
                "password": "siswa123",
                "portal": "admin"
            })
            assert resp.status_code == 403, f"Expected 403, got {resp.status_code}"
            self.log("Siswa correctly rejected at admin portal")
        self.test("Siswa login rejected at admin portal", _test)

    # ===== DASHBOARD TESTS =====
    def test_dashboard_ringkasan(self):
        """Test admin dashboard loads counts"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/dashboard", headers=headers)
            assert resp.status_code == 200, f"Dashboard failed: {resp.status_code} {resp.text}"
            data = resp.json()
            required_fields = ["books", "titles", "available", "members", "borrowed", "missing", "transactions", "visitors_today", "visitors_inside"]
            for field in required_fields:
                assert field in data, f"Dashboard missing field: {field}"
            self.log(f"Dashboard: {data['books']} books, {data['members']} members, {data['transactions']} transactions")
        self.test("Admin dashboard Ringkasan loads counts", _test)

    # ===== BOOKS TESTS =====
    def test_books_list(self):
        """Test GET /api/books"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/books", headers=headers)
            assert resp.status_code == 200, f"Books list failed: {resp.status_code} {resp.text}"
            books = resp.json()
            assert isinstance(books, list), "Books should be a list"
            self.log(f"Found {len(books)} books")
        self.test("GET /api/books returns book list", _test)

    def test_add_book(self):
        """Test POST /api/books"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            book_data = {
                "title": "TEST_Book_Migration",
                "author": "TEST_Author",
                "category": "Test",
                "shelf": "R-01",
                "isbn": "TEST123",
                "stock": 1
            }
            resp = requests.post(f"{BASE_URL}/books", json=book_data, headers=headers)
            assert resp.status_code == 200, f"Add book failed: {resp.status_code} {resp.text}"
            book = resp.json()
            assert "id" in book, "Book should have an id"
            self.test_book_id = book["id"]
            self.log(f"Created book: {self.test_book_id}")
        self.test("POST /api/books adds new book", _test)

    def test_edit_book(self):
        """Test PUT /api/books/{id}"""
        def _test():
            assert self.test_book_id, "No test book available"
            headers = {"Cookie": f"access_token={self.admin_token}"}
            update_data = {
                "title": "TEST_Book_Updated",
                "author": "TEST_Author",
                "category": "Test",
                "shelf": "R-01",
                "isbn": "TEST123"
            }
            resp = requests.put(f"{BASE_URL}/books/{self.test_book_id}", json=update_data, headers=headers)
            assert resp.status_code == 200, f"Edit book failed: {resp.status_code} {resp.text}"
            book = resp.json()
            assert book["title"] == "TEST_Book_Updated", "Book title not updated"
            self.log("Book updated successfully")
        self.test("PUT /api/books/{id} edits book", _test)

    def test_adjust_stock(self):
        """Test PATCH /api/books/{id}/stock"""
        def _test():
            assert self.test_book_id, "No test book available"
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.patch(f"{BASE_URL}/books/{self.test_book_id}/stock", json={"stock": 5}, headers=headers)
            assert resp.status_code == 200, f"Stock adjust failed: {resp.status_code} {resp.text}"
            book = resp.json()
            assert book["stock"] == 5, f"Expected stock 5, got {book.get('stock')}"
            self.log("Stock adjusted to 5")
        self.test("PATCH /api/books/{id}/stock adjusts stock", _test)

    def test_report_missing_book(self):
        """Test PATCH /api/books/{id}/status (lapor hilang)"""
        def _test():
            assert self.test_book_id, "No test book available"
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.patch(f"{BASE_URL}/books/{self.test_book_id}/status", json={
                "status": "Hilang",
                "note": "Test missing report",
                "reporter": "Test Admin"
            }, headers=headers)
            assert resp.status_code == 200, f"Report missing failed: {resp.status_code} {resp.text}"
            self.log("Book reported missing")
        self.test("PATCH /api/books/{id}/status reports missing book", _test)

    def test_found_missing_book(self):
        """Test PATCH /api/books/{id}/status (ditemukan)"""
        def _test():
            assert self.test_book_id, "No test book available"
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.patch(f"{BASE_URL}/books/{self.test_book_id}/status", json={
                "status": "Tersedia"
            }, headers=headers)
            assert resp.status_code == 200, f"Found book failed: {resp.status_code} {resp.text}"
            self.log("Missing book marked as found")
        self.test("PATCH /api/books/{id}/status marks book as found", _test)

    # ===== MEMBERS TESTS =====
    def test_members_list(self):
        """Test GET /api/members"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/members", headers=headers)
            assert resp.status_code == 200, f"Members list failed: {resp.status_code} {resp.text}"
            members = resp.json()
            assert isinstance(members, list), "Members should be a list"
            self.log(f"Found {len(members)} members")
        self.test("GET /api/members returns member list", _test)

    def test_add_member(self):
        """Test POST /api/members"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            member_data = {
                "name": "TEST_Member_Migration",
                "student_id": "TEST999",
                "class_name": "XII IPA 1"
            }
            resp = requests.post(f"{BASE_URL}/members", json=member_data, headers=headers)
            assert resp.status_code == 200, f"Add member failed: {resp.status_code} {resp.text}"
            member = resp.json()
            assert "id" in member, "Member should have an id"
            self.test_member_id = member["id"]
            self.log(f"Created member: {self.test_member_id}")
        self.test("POST /api/members adds new member", _test)

    def test_import_template_download(self):
        """Test GET /api/members/import-template.xlsx"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/members/import-template.xlsx", headers=headers)
            assert resp.status_code == 200, f"Template download failed: {resp.status_code} {resp.text}"
            assert "spreadsheet" in resp.headers.get("content-type", ""), "Wrong content type"
            assert len(resp.content) > 1000, "Template file too small"
            self.log(f"Template downloaded: {len(resp.content)} bytes")
        self.test("GET /api/members/import-template.xlsx downloads template", _test)

    def test_member_account_creation(self):
        """Test PUT /api/members/{id}/account (lock icon)"""
        def _test():
            assert self.test_member_id, "No test member available"
            headers = {"Cookie": f"access_token={self.admin_token}"}
            account_data = {
                "email": "test_member@test.com",
                "password": "test123"
            }
            resp = requests.put(f"{BASE_URL}/members/{self.test_member_id}/account", json=account_data, headers=headers)
            assert resp.status_code == 200, f"Account creation failed: {resp.status_code} {resp.text}"
            account = resp.json()
            assert account["email"] == "test_member@test.com", "Account email not set"
            self.log("Member account created")
        self.test("PUT /api/members/{id}/account creates member account", _test)

    def test_student_photo_switch(self):
        """Test GET /api/settings/student-photo (Foto siswa switch)"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/settings/student-photo", headers=headers)
            assert resp.status_code == 200, f"Student photo settings failed: {resp.status_code} {resp.text}"
            settings = resp.json()
            assert "enabled" in settings, "Settings should have 'enabled' field"
            self.log(f"Student photo upload: {'enabled' if settings['enabled'] else 'disabled'}")
        self.test("GET /api/settings/student-photo returns photo upload settings", _test)

    # ===== PHOTO UPLOAD TESTS =====
    def test_photo_upload_endpoint(self):
        """Test POST /api/uploads/photo exists"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            # Just check the endpoint exists (will fail without actual file, but shouldn't 404)
            resp = requests.post(f"{BASE_URL}/uploads/photo", headers=headers)
            assert resp.status_code != 404, "Photo upload endpoint not found"
            self.log("Photo upload endpoint exists")
        self.test("POST /api/uploads/photo endpoint exists", _test)

    # ===== SHELVES TESTS =====
    def test_shelves_list(self):
        """Test GET /api/shelves"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/shelves", headers=headers)
            assert resp.status_code == 200, f"Shelves list failed: {resp.status_code} {resp.text}"
            shelves = resp.json()
            assert isinstance(shelves, list), "Shelves should be a list"
            assert len(shelves) > 0, "Should have at least one shelf"
            self.log(f"Found {len(shelves)} shelves")
        self.test("GET /api/shelves returns shelf list", _test)

    def test_add_shelf(self):
        """Test POST /api/shelves"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            shelf_data = {
                "id": "TEST-01",
                "name": "TEST Shelf Migration",
                "capacity": 50,
                "color": "azure"
            }
            resp = requests.post(f"{BASE_URL}/shelves", json=shelf_data, headers=headers)
            assert resp.status_code == 200, f"Add shelf failed: {resp.status_code} {resp.text}"
            shelf = resp.json()
            assert shelf["id"] == "TEST-01", "Shelf id not set correctly"
            self.test_shelf_id = shelf["id"]
            self.log(f"Created shelf: {self.test_shelf_id}")
        self.test("POST /api/shelves adds new shelf", _test)

    def test_stocktake_list(self):
        """Test GET /api/stocktake (stok opname)"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/stocktake", headers=headers)
            assert resp.status_code == 200, f"Stocktake list failed: {resp.status_code} {resp.text}"
            stocktakes = resp.json()
            assert isinstance(stocktakes, list), "Stocktakes should be a list"
            self.log(f"Found {len(stocktakes)} stocktake records")
        self.test("GET /api/stocktake returns stocktake list", _test)

    # ===== CIRCULATION TESTS =====
    def test_borrow_book(self):
        """Test POST /api/transactions (pinjam)"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            # Get a member and book first
            members_resp = requests.get(f"{BASE_URL}/members", headers=headers)
            members = members_resp.json()
            if not members:
                self.log("No members available, skipping borrow test")
                return
            
            books_resp = requests.get(f"{BASE_URL}/books", headers=headers)
            books = books_resp.json()
            available_books = [b for b in books if b.get("available", 0) > 0]
            if not available_books:
                self.log("No available books, skipping borrow test")
                return
            
            transaction_data = {
                "member_id": members[0]["id"],
                "book_id": available_books[0]["id"],
                "action": "pinjam"
            }
            resp = requests.post(f"{BASE_URL}/transactions", json=transaction_data, headers=headers)
            assert resp.status_code == 200, f"Borrow failed: {resp.status_code} {resp.text}"
            self.log("Book borrowed successfully")
        self.test("POST /api/transactions (pinjam) borrows book", _test)

    def test_return_book(self):
        """Test POST /api/transactions (kembali)"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            # Get transactions to find a borrowed book
            tx_resp = requests.get(f"{BASE_URL}/transactions", headers=headers)
            transactions = tx_resp.json()
            borrowed = [t for t in transactions if t.get("action") == "pinjam"]
            if not borrowed:
                self.log("No borrowed books, skipping return test")
                return
            
            last_borrow = borrowed[0]
            transaction_data = {
                "member_id": last_borrow["member_id"],
                "book_id": last_borrow["book_id"],
                "action": "kembali"
            }
            resp = requests.post(f"{BASE_URL}/transactions", json=transaction_data, headers=headers)
            assert resp.status_code == 200, f"Return failed: {resp.status_code} {resp.text}"
            self.log("Book returned successfully")
        self.test("POST /api/transactions (kembali) returns book", _test)

    # ===== VISITORS TESTS =====
    def test_visitor_checkin_manual(self):
        """Test POST /api/visits/scan (manual NIS entry)"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            # Get a member
            members_resp = requests.get(f"{BASE_URL}/members", headers=headers)
            members = members_resp.json()
            if not members:
                self.log("No members available, skipping visitor test")
                return
            
            visit_data = {
                "code": members[0]["student_id"] or members[0]["id"]
            }
            resp = requests.post(f"{BASE_URL}/visits/scan", json=visit_data, headers=headers)
            assert resp.status_code == 200, f"Visitor check-in failed: {resp.status_code} {resp.text}"
            result = resp.json()
            if result.get("visit"):
                self.test_visit_id = result["visit"]["id"]
            self.log(f"Visitor check-in: {result.get('action')}")
        self.test("POST /api/visits/scan checks in visitor", _test)

    def test_guest_entry(self):
        """Test POST /api/visits/guest (guest entry)"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            guest_data = {
                "name": "TEST Guest Migration",
                "origin": "Test School",
                "purpose": "Test visit"
            }
            resp = requests.post(f"{BASE_URL}/visits/guest", json=guest_data, headers=headers)
            assert resp.status_code == 200, f"Guest entry failed: {resp.status_code} {resp.text}"
            self.log("Guest entry successful")
        self.test("POST /api/visits/guest creates guest entry", _test)

    def test_manual_checkout(self):
        """Test DELETE /api/visits/{id} (manual checkout)"""
        def _test():
            if not self.test_visit_id:
                self.log("No test visit available, skipping checkout test")
                return
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.delete(f"{BASE_URL}/visits/{self.test_visit_id}", headers=headers)
            assert resp.status_code == 200, f"Manual checkout failed: {resp.status_code} {resp.text}"
            self.log("Manual checkout successful")
        self.test("DELETE /api/visits/{id} performs manual checkout", _test)

    # ===== REPORTS TESTS =====
    def test_reports_transaksi(self):
        """Test GET /api/reports (transaksi)"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/reports?type=transaksi&period=semua", headers=headers)
            assert resp.status_code == 200, f"Reports transaksi failed: {resp.status_code} {resp.text}"
            self.log("Reports transaksi loaded")
        self.test("GET /api/reports (transaksi) returns transaction reports", _test)

    def test_reports_pdf_export(self):
        """Test GET /api/reports/export.pdf"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/reports/export.pdf?type=transaksi&period=semua", headers=headers)
            assert resp.status_code == 200, f"PDF export failed: {resp.status_code} {resp.text}"
            assert resp.headers.get("content-type") == "application/pdf", "Wrong content type"
            self.log(f"PDF export successful: {len(resp.content)} bytes")
        self.test("GET /api/reports/export.pdf exports PDF", _test)

    def test_reports_excel_export(self):
        """Test GET /api/reports/export.xlsx"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/reports/export.xlsx?type=masuk&period=semua", headers=headers)
            assert resp.status_code == 200, f"Excel export failed: {resp.status_code} {resp.text}"
            assert "spreadsheet" in resp.headers.get("content-type", ""), "Wrong content type"
            self.log(f"Excel export successful: {len(resp.content)} bytes")
        self.test("GET /api/reports/export.xlsx exports Excel", _test)

    # ===== PRINT ORIENTATION TESTS =====
    def test_pdf_portrait_orientation(self):
        """Test PDF export with portrait orientation"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/reports/export.pdf?type=transaksi&period=semua&orientation=portrait", headers=headers)
            assert resp.status_code == 200, f"PDF portrait export failed: {resp.status_code} {resp.text}"
            assert resp.headers.get("content-type") == "application/pdf", "Wrong content type"
            # Check filename contains -potret
            content_disp = resp.headers.get("content-disposition", "")
            assert "-potret-" in content_disp, f"Filename should contain -potret, got: {content_disp}"
            # Verify PDF is valid and contains expected content
            assert resp.content.startswith(b"%PDF"), "Should be a valid PDF"
            assert len(resp.content) > 1000, "PDF should have content"
            self.log(f"PDF portrait: {len(resp.content)} bytes, filename contains -potret ✓")
        self.test("PDF export with orientation=portrait has correct filename", _test)

    def test_pdf_landscape_orientation(self):
        """Test PDF export with landscape orientation"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/reports/export.pdf?type=masuk&period=semua&orientation=landscape", headers=headers)
            assert resp.status_code == 200, f"PDF landscape export failed: {resp.status_code} {resp.text}"
            assert resp.headers.get("content-type") == "application/pdf", "Wrong content type"
            # Check filename contains -lanskap
            content_disp = resp.headers.get("content-disposition", "")
            assert "-lanskap-" in content_disp, f"Filename should contain -lanskap, got: {content_disp}"
            # Verify PDF is valid and contains expected content
            assert resp.content.startswith(b"%PDF"), "Should be a valid PDF"
            assert len(resp.content) > 1000, "PDF should have content"
            self.log(f"PDF landscape: {len(resp.content)} bytes, filename contains -lanskap ✓")
        self.test("PDF export with orientation=landscape has correct filename", _test)

    def test_pdf_default_orientation(self):
        """Test PDF export without orientation param defaults to landscape"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/reports/export.pdf?type=hilang&period=semua", headers=headers)
            assert resp.status_code == 200, f"PDF default export failed: {resp.status_code} {resp.text}"
            # Check filename contains -lanskap (default)
            content_disp = resp.headers.get("content-disposition", "")
            assert "-lanskap-" in content_disp, f"Default should be landscape, got: {content_disp}"
            self.log("PDF default orientation is landscape ✓")
        self.test("PDF export without orientation param defaults to landscape", _test)

    def test_excel_portrait_orientation(self):
        """Test Excel export with portrait orientation"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/reports/export.xlsx?type=kunjungan&period=semua&orientation=portrait", headers=headers)
            assert resp.status_code == 200, f"Excel portrait export failed: {resp.status_code} {resp.text}"
            assert "spreadsheet" in resp.headers.get("content-type", ""), "Wrong content type"
            # Check filename contains -potret
            content_disp = resp.headers.get("content-disposition", "")
            assert "-potret-" in content_disp, f"Filename should contain -potret, got: {content_disp}"
            # Verify Excel orientation setting
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(resp.content))
            ws = wb.active
            assert ws.page_setup.orientation == "portrait", f"Excel orientation should be portrait, got {ws.page_setup.orientation}"
            self.log("Excel portrait orientation set correctly ✓")
        self.test("Excel export with orientation=portrait has correct page setup", _test)

    def test_excel_landscape_orientation(self):
        """Test Excel export with landscape orientation"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/reports/export.xlsx?type=statistik&period=7hari&orientation=landscape", headers=headers)
            assert resp.status_code == 200, f"Excel landscape export failed: {resp.status_code} {resp.text}"
            assert "spreadsheet" in resp.headers.get("content-type", ""), "Wrong content type"
            # Check filename contains -lanskap
            content_disp = resp.headers.get("content-disposition", "")
            assert "-lanskap-" in content_disp, f"Filename should contain -lanskap, got: {content_disp}"
            # Verify Excel orientation setting
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(resp.content))
            ws = wb.active
            assert ws.page_setup.orientation == "landscape", f"Excel orientation should be landscape, got {ws.page_setup.orientation}"
            self.log("Excel landscape orientation set correctly ✓")
        self.test("Excel export with orientation=landscape has correct page setup", _test)

    def test_all_report_types_with_orientation(self):
        """Test all report types support orientation parameter"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            report_types = ["transaksi", "masuk", "hilang", "kunjungan", "statistik", "peringkat"]
            for report_type in report_types:
                for orientation in ["portrait", "landscape"]:
                    resp = requests.get(f"{BASE_URL}/reports/export.pdf?type={report_type}&period=semua&orientation={orientation}", headers=headers)
                    assert resp.status_code == 200, f"PDF {report_type} {orientation} failed: {resp.status_code}"
                    content_disp = resp.headers.get("content-disposition", "")
                    expected = "-potret-" if orientation == "portrait" else "-lanskap-"
                    assert expected in content_disp, f"{report_type} {orientation} filename wrong: {content_disp}"
            self.log(f"All {len(report_types)} report types support both orientations ✓")
        self.test("All report types (transaksi, masuk, hilang, kunjungan, statistik, peringkat) support orientation", _test)

    def test_backup_zip_not_affected(self):
        """Test backup.zip endpoint still works (no orientation param)"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/admin/backup.zip", headers=headers)
            assert resp.status_code == 200, f"Backup ZIP failed: {resp.status_code} {resp.text}"
            assert resp.headers.get("content-type") == "application/zip", "Wrong content type"
            # Verify it's a valid ZIP
            try:
                z = zipfile.ZipFile(io.BytesIO(resp.content))
                files = z.namelist()
                assert "data-cadangan.json" in files, "Backup should contain data-cadangan.json"
                self.log(f"Backup ZIP contains {len(files)} files ✓")
            except zipfile.BadZipFile:
                raise AssertionError("Invalid ZIP file")
        self.test("GET /api/admin/backup.zip still works (backward compatible)", _test)

    def test_reports_summary(self):
        """Test GET /api/reports/summary (Ringkasan & grafik)"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/reports/summary", headers=headers)
            assert resp.status_code == 200, f"Reports summary failed: {resp.status_code} {resp.text}"
            self.log("Reports summary loaded")
        self.test("GET /api/reports/summary returns summary and graphs", _test)

    def test_visits_leaderboard(self):
        """Test GET /api/visits/leaderboard (Peringkat)"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/visits/leaderboard", headers=headers)
            assert resp.status_code == 200, f"Leaderboard failed: {resp.status_code} {resp.text}"
            self.log("Leaderboard loaded")
        self.test("GET /api/visits/leaderboard returns leaderboard", _test)

    # ===== BACKUP TESTS =====
    def test_backup_zip(self):
        """Test GET /api/admin/backup.zip"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.get(f"{BASE_URL}/admin/backup.zip", headers=headers)
            assert resp.status_code == 200, f"Backup failed: {resp.status_code} {resp.text}"
            assert resp.headers.get("content-type") == "application/zip", "Wrong content type"
            self.log(f"Backup ZIP downloaded: {len(resp.content)} bytes")
        self.test("GET /api/admin/backup.zip works for admin", _test)

    def test_reset_code_email(self):
        """Test POST /api/admin/reset-code?purpose=restore (email verification)"""
        def _test():
            headers = {"Cookie": f"access_token={self.admin_token}"}
            resp = requests.post(f"{BASE_URL}/admin/reset-code?purpose=restore", headers=headers)
            assert resp.status_code == 200, f"Reset code failed: {resp.status_code} {resp.text}"
            result = resp.json()
            assert result.get("ok") is True, "Reset code should return ok: true"
            assert "sent_to" in result, "Should return masked email"
            self.log(f"Reset code sent to: {result['sent_to']}")
        self.test("POST /api/admin/reset-code?purpose=restore sends email (200 response)", _test)

    # ===== SISWA PORTAL TESTS =====
    def test_siswa_catalog(self):
        """Test siswa can access catalog"""
        def _test():
            headers = {"Cookie": f"access_token={self.siswa_token}"}
            resp = requests.get(f"{BASE_URL}/books", headers=headers)
            assert resp.status_code == 200, f"Siswa catalog failed: {resp.status_code} {resp.text}"
            books = resp.json()
            self.log(f"Siswa can see {len(books)} books")
        self.test("Siswa can access catalog at /siswa portal", _test)

    def test_siswa_shelves(self):
        """Test siswa can access rak"""
        def _test():
            headers = {"Cookie": f"access_token={self.siswa_token}"}
            resp = requests.get(f"{BASE_URL}/shelves", headers=headers)
            assert resp.status_code == 200, f"Siswa shelves failed: {resp.status_code} {resp.text}"
            shelves = resp.json()
            self.log(f"Siswa can see {len(shelves)} shelves")
        self.test("Siswa can access rak at /siswa portal", _test)

    def test_siswa_reservations(self):
        """Test siswa can access reservations (Pesanan saya)"""
        def _test():
            headers = {"Cookie": f"access_token={self.siswa_token}"}
            resp = requests.get(f"{BASE_URL}/reservations/mine", headers=headers)
            assert resp.status_code == 200, f"Siswa reservations failed: {resp.status_code} {resp.text}"
            reservations = resp.json()
            self.log(f"Siswa has {len(reservations)} reservations")
        self.test("Siswa can access reservations (Pesanan saya)", _test)

    def test_siswa_leaderboard(self):
        """Test siswa can access Peringkat tab"""
        def _test():
            headers = {"Cookie": f"access_token={self.siswa_token}"}
            resp = requests.get(f"{BASE_URL}/visits/leaderboard", headers=headers)
            assert resp.status_code == 200, f"Siswa leaderboard failed: {resp.status_code} {resp.text}"
            self.log("Siswa can access Peringkat tab")
        self.test("Siswa can access Peringkat tab", _test)

    # ===== CLEANUP =====
    def cleanup(self):
        """Delete test data"""
        print(f"\n🧹 Cleaning up test data...")
        headers = {"Cookie": f"access_token={self.admin_token}"}
        
        # Delete test book
        if self.test_book_id:
            try:
                resp = requests.delete(f"{BASE_URL}/books/{self.test_book_id}", headers=headers)
                if resp.status_code == 200:
                    print(f"  ✓ Deleted book {self.test_book_id}")
            except Exception as e:
                print(f"  ✗ Error deleting book: {e}")
        
        # Delete test member
        if self.test_member_id:
            try:
                resp = requests.delete(f"{BASE_URL}/members/{self.test_member_id}", headers=headers)
                if resp.status_code == 200:
                    print(f"  ✓ Deleted member {self.test_member_id}")
            except Exception as e:
                print(f"  ✗ Error deleting member: {e}")
        
        # Delete test shelf
        if self.test_shelf_id:
            try:
                resp = requests.delete(f"{BASE_URL}/shelves/{self.test_shelf_id}", headers=headers)
                if resp.status_code == 200:
                    print(f"  ✓ Deleted shelf {self.test_shelf_id}")
            except Exception as e:
                print(f"  ✗ Error deleting shelf: {e}")

    def summary(self):
        print(f"\n{'='*60}")
        print(f"📊 Test Summary")
        print(f"{'='*60}")
        print(f"Total tests: {self.tests_run}")
        print(f"Passed: {self.tests_passed}")
        print(f"Failed: {self.tests_run - self.tests_passed}")
        print(f"Success rate: {self.tests_passed / self.tests_run * 100:.1f}%")
        print(f"{'='*60}")
        return 0 if self.tests_passed == self.tests_run else 1

def main():
    runner = TestRunner()
    
    try:
        # Auth tests
        runner.test_admin_login()
        runner.test_petugas_login()
        runner.test_siswa_login_siswa_portal()
        runner.test_siswa_rejected_at_admin_portal()
        
        # Dashboard
        runner.test_dashboard_ringkasan()
        
        # Books
        runner.test_books_list()
        runner.test_add_book()
        runner.test_edit_book()
        runner.test_adjust_stock()
        runner.test_report_missing_book()
        runner.test_found_missing_book()
        
        # Members
        runner.test_members_list()
        runner.test_add_member()
        runner.test_import_template_download()
        runner.test_member_account_creation()
        runner.test_student_photo_switch()
        
        # Photo upload
        runner.test_photo_upload_endpoint()
        
        # Shelves
        runner.test_shelves_list()
        runner.test_add_shelf()
        runner.test_stocktake_list()
        
        # Circulation
        runner.test_borrow_book()
        runner.test_return_book()
        
        # Visitors
        runner.test_visitor_checkin_manual()
        runner.test_guest_entry()
        runner.test_manual_checkout()
        
        # Reports
        runner.test_reports_transaksi()
        runner.test_reports_pdf_export()
        runner.test_reports_excel_export()
        runner.test_reports_summary()
        runner.test_visits_leaderboard()
        
        # Print orientation tests
        runner.test_pdf_portrait_orientation()
        runner.test_pdf_landscape_orientation()
        runner.test_pdf_default_orientation()
        runner.test_excel_portrait_orientation()
        runner.test_excel_landscape_orientation()
        runner.test_all_report_types_with_orientation()
        runner.test_backup_zip_not_affected()
        
        # Backup
        runner.test_backup_zip()
        runner.test_reset_code_email()
        
        # Siswa portal
        runner.test_siswa_catalog()
        runner.test_siswa_shelves()
        runner.test_siswa_reservations()
        runner.test_siswa_leaderboard()
        
    finally:
        # Always cleanup
        runner.cleanup()
    
    return runner.summary()

if __name__ == "__main__":
    sys.exit(main())
